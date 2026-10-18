"""Create a fourteen-slide beginner presentation and a matching PDF.

Stage the files in /tmp before copying the reviewed files to Windows Desktop.
"""
import math
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR
from PIL import Image, ImageDraw
from generate_networking_presentation import (
    add_text, card, line, circle, arrow, set_bg, footer,
    INK, SLATE, WHITE, PAPER, NAVY, CYAN_DARK, CYAN_PALE,
    MINT_DARK, MINT_PALE, PURPLE_DARK, PURPLE_PALE, PP_ALIGN,
)
from create_mobile_safe_presentation import SlideRenderer, color

OUT = Path("/tmp/networking_easy_14")
OUT.mkdir(exist_ok=True)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

def text(s, t, x, y, w, h=.65, size=22, color=INK, bold=False, center=False):
    # Use paragraph spacing instead of blank paragraphs, so every line has
    # explicit font settings in PowerPoint and the PDF renderer.
    separated = "\n\n" in t
    t = t.replace("\n\n", "\n")
    shape = add_text(s, t, x, y, w, h, size=size, color=color, bold=bold,
                    font="DejaVu Sans", margin=0,
                    align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT)
    for p in shape.text_frame.paragraphs[:-1]:
        p.space_after = Pt(8 if separated else 3)
    return shape

def slide(title, subtitle, section):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, WHITE)
    text(s, section.upper(), .65, .3, 12, .3, 11, CYAN_DARK, True)
    text(s, title, .65, .85, 12, .65, 30, bold=True)
    text(s, subtitle, .65, 1.58, 12, .6, 17, SLATE)
    footer(s, len(prs.slides), section)
    return s

def panel(s, title, body, x, y, w=3.85, h=3.4, fill=CYAN_PALE, accent=CYAN_DARK):
    card(s, x, y, w, h, fill, fill)
    text(s, title, x+.24, y+.24, w-.48, .8, 23, accent, True)
    text(s, body, x+.24, y+.94, w-.48, h-1.08, 19)

def note(s, t):
    card(s, .65, 6.25, 12, .62, PAPER, PAPER)
    text(s, t, .85, 6.38, 11.6, .4, 17, CYAN_DARK, True)

def node(s, name, x, y, fill=CYAN_PALE):
    circle(s, x, y, .7, fill, CYAN_DARK, 2)
    text(s, name, x, y+.18, .7, .4, 16, bold=True, center=True)

s=slide("How computer networks move data", "IPv4 to IPv6 strategies • Packet network topology • Datagrams and virtual circuits", "Computer networks | Easy guide")
panel(s,"01  Network layout", "See how devices and routers connect.\n\nExplore six network layouts with diagrams.",.65,2.65)
panel(s,"02  Packet delivery", "Follow a packet's journey.\n\nCompare datagrams with virtual circuits.",4.73,2.65,fill=MINT_PALE,accent=MINT_DARK)
panel(s,"03  IPv4 and IPv6", "Understand why IPv6 is needed.\n\nExplore dual stack, tunnelling, and translation.",8.81,2.65,fill=PURPLE_PALE,accent=PURPLE_DARK)
note(s,"14 slides • Simple definitions • Network examples • Clear diagrams")

s=slide("What is a packet network?", "A large message travels as smaller pieces called packets.","01 | Packet networks")
panel(s,"Inside a packet", "Header: delivery details, including source and destination IP addresses.\n\nPayload: the data being carried, such as part of a photo or message.",.65,2.45,w=5.3,h=3.4)
text(s,"Example: sending a photo",6.55,2.55,5.5,.5,23,bold=True)
for i in range(3):
    card(s,6.65+i*1.85,3.5,1.5,.8,MINT_PALE,MINT_DARK)
    text(s,f"Part {i+1}",6.65+i*1.85,3.72,1.5,.4,20,center=True)
text(s,"Small packets share links with other users' traffic.\nThe receiver puts the data back together.",6.55,4.65,5.7,1.4,20)
note(s,"Packet switching means forwarding small packets through shared network links.")

s=slide("A packet's journey: sender to receiver", "Example: your laptop sends a photo to a friend's device.","02 | Step-by-step journey")
journey=[("Laptop", "Makes packets"), ("Router A", "Chooses next hop"), ("Router B", "Forwards onward"), ("Friend", "Rebuilds the photo")]
for i,(head,body) in enumerate(journey):
    x=.65+i*3.08
    card(s,x,2.45,2.76,1.35,CYAN_PALE if i in (0,3) else MINT_PALE)
    text(s,head,x+.15,2.66,2.46,.45,22,bold=True,center=True)
    text(s,body,x+.1,3.22,2.56,.4,16,center=True)
    if i<3:arrow(s,x+2.82,2.98,.2,.3,CYAN_DARK)
text(s,"What happens at each router?",.8,4.12,11.7,.5,23,CYAN_DARK,True)
text(s,"1. Read the destination IP address.\n2. Check a routing table: a list of where to send packets next.\n3. Forward the packet through the selected link.",.8,4.78,11.7,1.25,20)
note(s,"If a link is busy, packets may wait in a queue. If the queue is full, some may be dropped.")
s.notes_slide.notes_text_frame.text="This is a simplified IP forwarding example. Routing protocols build forwarding information; each packet is sent to a next hop. End systems use transport and application protocols to reconstruct the original data. A dropped packet is retransmitted only if a suitable protocol, such as TCP, provides that behaviour."

s=slide("Topology means the network's layout", "Devices are the points; links are the connections between them.","03 | Packet network topology")
for a,b,c,d in [(1.65,3.9,3.15,3.05),(1.65,3.9,3.15,4.75),(3.15,3.05,4.85,3.9),(3.15,4.75,4.85,3.9)]:
    line(s,a,b,c,d,CYAN_DARK,3)
for name,x,y in [("A",1.3,3.55),("B",2.8,2.7),("C",2.8,4.4),("D",4.5,3.55)]:node(s,name,x,y)
text(s,"Two possible routes: A–B–D or A–C–D",.85,5.55,5.6,.6,18,center=True)
panel(s,"Three useful terms", "Node: a connected device or router.\n\nLink: a wired or wireless connection.\n\nRoute: the path taken through the network.",6.65,2.45,w=6,h=3.5,fill=MINT_PALE,accent=MINT_DARK)
note(s,"If the A–B link fails, A–C–D offers another route once routing has adjusted.")

s=slide("Common network topologies", "The layout affects cost, speed, and how a network handles failures.","04 | Topology examples")
for idx,(name,desc) in enumerate([("Star","All devices connect to a central device. Easy to manage; the centre is a key failure point."),("Ring","Devices form a loop. Traffic follows the ring; a simple ring can be disrupted by a break."),("Mesh","Devices have several connections. Alternative paths improve resilience but cost more.")]):
    x=.65+idx*4.08
    card(s,x,2.4,3.85,3.65,PAPER,PAPER)
    text(s,name,x+.2,2.6,3.45,.45,24,CYAN_DARK,True)
    points=[(x+.75,3.25),(x+2.5,3.25),(x+2.5,4.1),(x+.75,4.1)]
    if idx==0:
        for px,py in points:line(s,px+.1,py+.1,x+1.72,3.78,CYAN_DARK,2)
        circle(s,x+1.6,3.66,.24,MINT_DARK)
    else:
        for j in range(4):
            p,q=points[j],points[(j+1)%4];line(s,p[0]+.1,p[1]+.1,q[0]+.1,q[1]+.1,CYAN_DARK,2)
        if idx==2:
            for a,b in [(0,2),(1,3)]:line(s,points[a][0]+.1,points[a][1]+.1,points[b][0]+.1,points[b][1]+.1,CYAN_DARK,2)
    for px,py in points:circle(s,px,py,.2,CYAN_DARK)
    text(s,desc,x+.2,4.65,3.45,1.2,18)
note(s,"Example: an office may use a star, while its provider uses a mesh of routers.")

s=slide("More topologies: bus, tree, and hybrid", "These layouts show other ways to connect devices and groups of devices.","05 | More network topologies")
for idx,(name,desc) in enumerate([
    ("Bus", "Devices share one main cable. It uses little cable, but a break in the main cable can disrupt the network."),
    ("Tree", "Connections branch into levels. It is easy to expand, but an upper-level failure can disconnect a whole branch."),
    ("Hybrid", "Combines two or more layouts. For example, star networks in separate rooms can connect through a ring.")
]):
    x=.65+idx*4.08
    card(s,x,2.4,3.85,3.65,PAPER,PAPER)
    text(s,name,x+.2,2.6,3.45,.45,24,CYAN_DARK,True)
    if idx==0:
        line(s,x+.45,3.87,x+3.4,3.87,CYAN_DARK,3)
        for j in range(4):
            px=x+.7+j*.8
            line(s,px,3.87,px,3.35,CYAN_DARK,2)
            circle(s,px-.1,3.23,.2,CYAN_DARK)
        for px in [x+.45,x+3.4]:line(s,px,3.69,px,4.05,CYAN_DARK,2)
    elif idx==1:
        pts=[(x+1.92,3.2),(x+1.02,3.68),(x+2.82,3.68),
             (x+.57,4.15),(x+1.47,4.15),(x+2.37,4.15),(x+3.27,4.15)]
        for a,b in [(0,1),(0,2),(1,3),(1,4),(2,5),(2,6)]:
            line(s,*pts[a],*pts[b],CYAN_DARK,2)
        for px,py in pts:circle(s,px-.1,py-.1,.2,CYAN_DARK)
    else:
        hubs=[(x+.85,3.58),(x+2.92,3.58),(x+1.9,4.2)]
        for j in range(3):line(s,*hubs[j],*hubs[(j+1)%3],CYAN_DARK,2)
        for j,(px,py) in enumerate(hubs):
            offsets=[(-.35,-.38),(.35,-.38)] if j<2 else [(-.48,.13),(.48,.13)]
            for dx,dy in offsets:
                line(s,px,py,px+dx,py+dy,CYAN_DARK,1.6)
                circle(s,px+dx-.07,py+dy-.07,.14,CYAN_DARK)
            circle(s,px-.1,py-.1,.2,MINT_DARK)
    text(s,desc,x+.2,4.65,3.45,1.2,18)
note(s,"Choose a layout based on cost, room to grow, and the effect of a failed device or link.")
s.notes_slide.notes_text_frame.text="Bus describes the classic shared-cable topology. Tree describes a hierarchy of branches. Hybrid combines different topologies; the diagram uses small stars attached through a ring. Real networks may include redundant links to reduce the effect of failures."

s=slide("Datagrams: each packet travels independently", "A datagram network does not set up a path before sending packets.","06 | Datagrams")
panel(s,"How it works", "Each packet carries a destination address.\n\nEach router chooses the next hop. Packets may take different routes.\n\nExample: the Internet Protocol (IP).",.65,2.4,w=5.7,h=3.6)
panel(s,"What this means", "A webpage or video call travels in IP packets.\n\nBenefit: sending can start without setting up a circuit.\n\nLimit: packets can be lost, delayed, or arrive out of order.",6.6,2.4,w=6.05,h=3.6,fill=MINT_PALE,accent=MINT_DARK)
note(s,"TCP, a protocol used above IP, can resend lost data and put it in order.")

s=slide("Virtual circuits: set up a path first", "A virtual circuit is a logical connection through a packet network.","07 | Virtual circuits")
for i,(head,body) in enumerate([("1. Set up","Choose a path first.\n\nSwitches store circuit labels and where to forward packets."),("2. Send","Packets follow this path using short labels.\n\nMany circuits can share the same physical link."),("3. Finish","A switched circuit is removed when the session ends.\n\nA permanent circuit stays configured.")]):
    panel(s,head,body,.65+i*4.08,2.5,h=3.5,fill=MINT_PALE,accent=MINT_DARK)
note(s,"An agreed route helps control forwarding; bandwidth and reliable delivery need extra support.")
s.notes_slide.notes_text_frame.text="Traditional examples include X.25 and ATM. Permanent virtual circuits are provisioned in advance and remain available; setup and teardown describe switched virtual circuits. A virtual circuit alone does not guarantee reliable delivery."

s=slide("A closer look: how packets are forwarded", "Devices use stored information to decide which link a packet should take next.","08 | More on datagrams and virtual circuits")
for x,head,body,flow,fill,accent in [
    (.65,"Datagram: address lookup", "A router reads the destination IP address and checks its routing table.\n\nEvery packet gets a lookup. Packets often take the same route, but do not have to.","Destination IP  →  Router  →  Next hop",CYAN_PALE,CYAN_DARK),
    (6.8,"Virtual circuit: label lookup", "A switch uses the incoming link and circuit label to find the outgoing link.\n\nThe circuit label can change at each hop while the established path stays the same.","Label 12  →  Switch  →  Label 7",MINT_PALE,MINT_DARK),
]:
    card(s,x,2.4,5.85,3.65,fill,fill)
    text(s,head,x+.24,2.66,5.37,.6,22,accent,True)
    text(s,body,x+.24,3.42,5.37,1.9,19)
    text(s,flow,x+.24,5.55,5.37,.4,17,accent,True,True)
note(s,"Both methods can face busy links. A virtual circuit does not automatically prevent packet loss.")
s.notes_slide.notes_text_frame.text="Circuit labels are local to links; actual label formats depend on the technology. The values 12 and 7 are illustrative circuit labels, not a specific ATM VPI/VCI configuration. Source: Cisco ATM command reference, https://www.cisco.com/c/en/us/td/docs/ios-xml/ios/atm/command/atm-cr-book/atm-m1.html . IP routers forward by destination-address lookup; an implementation may use an optimised forwarding table derived from routing information."

s=slide("Datagrams and virtual circuits compared", "Both carry packets, but they organise forwarding differently.","09 | Quick comparison")
rows=[("Feature","Datagram","Virtual circuit"),("Path setup","No setup required","Path established first"),("Route","May differ by packet","Same established path"),("Packet identifier","Destination address","Short circuit label"),("Network state","No per-circuit state needed","Stores circuit information"),("Path failure","Routing can find a new route","Circuit may need rebuilding"),("Example","IP networks","ATM and X.25")]
for r,row in enumerate(rows):
    y=2.35+r*.52
    card(s,.65,y,12,.49,CYAN_PALE if r==0 else PAPER,CYAN_PALE if r==0 else PAPER)
    for c,t in enumerate(row):text(s,t,.85+c*4,y+.07,3.8,.4,16,CYAN_DARK if r==0 else INK,r==0)
note(s,"A virtual circuit is a logical path; it does not require a dedicated physical wire.")

s=slide("Why move from IPv4 to IPv6?", "An IP address identifies a device's network interface for communication.","10 | IPv4 and IPv6")
panel(s,"IPv4", "32-bit addresses: about 4.3 billion possibilities.\n\nExample: 192.0.2.10\n\nLimited addresses led to NAT: several devices share one public IPv4 address.",.65,2.45,w=5.85,h=3.65)
panel(s,"IPv6", "128-bit addresses: a vastly larger address space.\n\nExample: 2001:db8::10\n\nProvides room for more devices. Both versions will coexist during migration.",6.8,2.45,w=5.85,h=3.65,fill=PURPLE_PALE,accent=PURPLE_DARK)
note(s,"IPv4 and IPv6 use different packet formats, so mixed networks need a transition method.")

s=slide("Dual stack and tunnelling in more detail", "Both methods help networks move toward IPv6 in stages.","11 | Transition strategies in detail")
for x,head,fill,accent in [(.65,"Dual stack: support both",CYAN_PALE,CYAN_DARK),(6.8,"Tunnelling: carry inside",MINT_PALE,MINT_DARK)]:
    card(s,x,2.4,5.85,3.72,fill,fill)
    text(s,head,x+.24,2.63,5.37,.55,23,accent,True)
text(s,"IPv4 website  ←  Device  →  IPv6 website",.89,3.38,5.37,.4,17,CYAN_DARK,True,True)
text(s,"The device and network support both versions. Each connection uses a suitable IP version.\n\nBenefit: devices can migrate gradually.\n\nTrade-off: two versions to manage; IPv4 addresses are still needed.",.89,4.02,5.37,1.92,18)
card(s,7.1,3.35,5.25,.52,WHITE,MINT_DARK)
text(s,"IPv4 outer packet",7.23,3.48,2.3,.3,15,MINT_DARK,True)
card(s,9.75,3.43,2.42,.36,MINT_PALE,MINT_PALE)
text(s,"IPv6 packet inside",9.8,3.49,2.32,.3,14,MINT_DARK,True,True)
text(s,"One tunnel endpoint wraps the IPv6 packet in IPv4. The other removes the outer wrapper.\n\nBenefit: links IPv6 sites through IPv4.\n\nTrade-off: extra header data and setup. A tunnel alone does not encrypt data.",7.04,4.02,5.37,1.92,18)
text(s,"Reference: RFC 4213, Basic Transition Mechanisms for IPv6 Hosts and Routers.",.8,6.65,12,.3,11,SLATE)
s.notes_slide.notes_text_frame.text="Source: https://www.rfc-editor.org/rfc/rfc4213.html . Diagram shows one example: IPv6 over IPv4. Both tunnel endpoints understand the inner protocol. Encapsulation adds an outer header; it is not encryption. Dual stack enables separate IPv4 and IPv6 communication and does not itself translate between IPv4-only and IPv6-only endpoints."

s=slide("Translation: what the gateway does", "Translation lets an IPv6-only side communicate with an IPv4-only side.","12 | Translation in detail")
for i,(head,body) in enumerate([
    ("1. Read", "Check the incoming packet's IP addresses and protocol information."),
    ("2. Convert", "Build the other IP version's header and map the addresses to suitable values."),
    ("3. Forward", "Send the converted packet onward. Replies are translated for the return journey.")
]):
    panel(s,head,body,.65+i*4.08,2.4,h=2.42,fill=PURPLE_PALE,accent=PURPLE_DARK)
text(s,"Useful because",.8,5.08,5.5,.45,22,PURPLE_DARK,True)
text(s,"An IPv6-only network can still reach services that only support IPv4.",.8,5.67,5.5,.8,19)
text(s,"What needs care",6.8,5.08,5.5,.45,22,PURPLE_DARK,True)
text(s,"NAT64 keeps a record of connections. Apps using IPv4 addresses directly may need extra help.",6.8,5.67,5.5,.8,19)
text(s,"Reference: RFC 6146 (stateful NAT64). NAT64 is one example of IPv4/IPv6 translation.",.8,6.65,12,.3,11,SLATE)
s.notes_slide.notes_text_frame.text="Source: https://www.rfc-editor.org/rfc/rfc6146.html . This slide's connection-tracking statement refers to stateful NAT64, not to all translation techniques. A translator maps IP headers and addresses and adjusts relevant transport checksums; the application data is normally preserved. An IPv4 literal used by an application can bypass DNS64 name synthesis, so NAT64 plus DNS64 alone may not meet that application's needs."

s=prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s,WHITE)
text(s,"COMPUTER NETWORKS",.65,.3,12,.3,11,CYAN_DARK,True)
line(s,5.2,2.35,8.1,2.35,CYAN_DARK,2)
for x,fill in [(5.05,CYAN_PALE),(6.5,MINT_PALE),(7.95,PURPLE_PALE)]:
    circle(s,x,2.2,.3,fill,CYAN_DARK,1.5)
text(s,"Thank You",.65,3.1,12,1.1,48,CYAN_DARK,True,True)
footer(s,len(prs.slides),"Thank you")

prs.core_properties.title="Computer Networks: IPv4–IPv6, Topology, Datagrams and Virtual Circuits"
prs.core_properties.subject="A fourteen-slide presentation in easy-to-understand language"
prs.save(OUT/"computer_networks_14_slides.pptx")

class SpacedRenderer(SlideRenderer):
    """Respect paragraph spacing and report text that exceeds its text box."""
    def render_text(self, draw, shape):
        if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
            return
        x1,y1,x2,y2=self.box(shape)
        tf=shape.text_frame
        left=x1+round(tf.margin_left*self.sx)
        right=x2-round(tf.margin_right*self.sx)
        top=y1+round(tf.margin_top*self.sy)
        bottom=y2-round(tf.margin_bottom*self.sy)
        lines=[]
        scale=self.pixels_per_inch/72
        for p in tf.paragraphs:
            if not p.runs:
                continue
            run=p.runs[0]
            font=self.font(run,p,run.font.size.pt if run.font.size else 12)
            ascent,descent=font.getmetrics()
            height=round(ascent+descent)
            wrapped=self.wrap(draw,p.text,font,right-left)
            for j,value in enumerate(wrapped):
                gap=round(p.space_after.pt*scale) if p.space_after and j==len(wrapped)-1 else 0
                lines.append((value,font,color(run.font.color),p.alignment,height,gap))
        total=sum(item[4]+item[5] for item in lines)
        if total>bottom-top+3:
            print(f"TEXT OVERFLOW slide {self.current_slide}: {shape.text[:75]!r} needs {total}px; has {bottom-top}px")
        y=top
        if tf.vertical_anchor==MSO_ANCHOR.MIDDLE:
            y+=max(0,(bottom-top-total)//2)
        elif tf.vertical_anchor==MSO_ANCHOR.BOTTOM:
            y+=max(0,bottom-top-total)
        for value,font,fill,align,height,gap in lines:
            width=draw.textlength(value,font=font)
            x=left+(right-left-width)/2 if align==PP_ALIGN.CENTER else right-width if align==PP_ALIGN.RIGHT else left
            draw.text((round(x),y),value,font=font,fill=fill,anchor="lt")
            y+=height+gap

renderer=SpacedRenderer(prs)
images=[]
for i,s in enumerate(prs.slides,1):
    renderer.current_slide=i
    im=renderer.render(s)
    im.save(OUT/f"slide_{i:02}.png")
    images.append(im)
images[0].save(OUT/"computer_networks_14_pages.pdf",save_all=True,append_images=images[1:],resolution=192)
thumb=Image.new("RGB",(1280,math.ceil(len(images)/2)*360),"#dddddd")
for i,im in enumerate(images):
    im=im.resize((640,360))
    thumb.paste(im,((i%2)*640,(i//2)*360))
thumb.save(OUT/"preview.jpg")
print(OUT)
