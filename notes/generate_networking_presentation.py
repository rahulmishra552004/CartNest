"""Generate a visual 10-slide deck about packet networks and IPv4/IPv6 transition.

Run with:
    PYTHONPATH=/tmp/network_pptx_deps python3 notes/generate_networking_presentation.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent / "packet_networks_ipv4_ipv6_translation.pptx"

W = 13.333
H = 7.5
FONT = "Aptos"
FONT_DISPLAY = "Aptos Display"

# Theme
NAVY = "0B1020"
NAVY_2 = "111A31"
INK = "172033"
SLATE = "46536B"
MUTED = "6B7890"
PAPER = "F7F8FC"
WHITE = "FFFFFF"
LINE = "DDE3EE"
CYAN = "35C2FF"
CYAN_DARK = "087EA4"
CYAN_PALE = "DDF6FF"
MINT = "2DD4BF"
MINT_DARK = "0F766E"
MINT_PALE = "DDFBF6"
PURPLE = "8B5CF6"
PURPLE_DARK = "6039B5"
PURPLE_PALE = "EEE8FF"
AMBER = "FBBF24"
AMBER_DARK = "A96500"
AMBER_PALE = "FFF3CC"
CORAL = "FB7185"
CORAL_DARK = "B73751"
CORAL_PALE = "FFE5EA"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_bg(slide, color: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=INK,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.04,
    line_spacing=1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    lines = str(text).split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, runs, x, y, w, h, size=16, color=INK, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    for item in runs:
        if isinstance(item, str):
            item = {"text": item}
        r = p.add_run()
        r.text = item["text"]
        r.font.name = item.get("font", FONT)
        r.font.size = Pt(item.get("size", size))
        r.font.bold = item.get("bold", False)
        r.font.color.rgb = rgb(item.get("color", color))
    return box


def add_shape(slide, kind, x, y, w, h, fill, line=None, radius=True, line_width=1.0):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(line_width)
    return shape


def card(slide, x, y, w, h, fill=WHITE, line=LINE, line_width=1.0):
    return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, line_width=line_width)


def pill(slide, text, x, y, w, fill, text_color=WHITE, size=10, line=None):
    card(slide, x, y, w, 0.34, fill, line or fill)
    add_text(slide, text, x + 0.04, y + 0.02, w - 0.08, 0.27, size=size, color=text_color, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def circle(slide, x, y, d, fill, line=None, line_width=1.0):
    return add_shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill, line or fill, line_width=line_width)


def line(slide, x1, y1, x2, y2, color=LINE, width=2.0, dash=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width)
    if dash:
        conn.line.dash_style = dash
    return conn


def arrow(slide, x, y, w, h, color=CYAN):
    return add_shape(slide, MSO_SHAPE.CHEVRON, x, y, w, h, color, color)


def title_block(slide, kicker, title, subtitle, dark=False, title_size=28):
    fg = WHITE if dark else INK
    sub = "B7C4DD" if dark else MUTED
    add_text(slide, kicker.upper(), 0.72, 0.38, 8.6, 0.25, size=10, color=CYAN, bold=True, margin=0)
    add_text(slide, title, 0.72, 0.70, 11.9, 0.54, size=title_size, color=fg, bold=True,
             font=FONT_DISPLAY, margin=0)
    if subtitle:
        add_text(slide, subtitle, 0.72, 1.27, 11.7, 0.40, size=13, color=sub, margin=0)


def footer(slide, number, section, dark=False):
    c = "42506A" if dark else "D8DFEA"
    t = "8FA0BF" if dark else MUTED
    line(slide, 0.72, 7.13, 12.61, 7.13, c, 0.8)
    add_text(slide, section.upper(), 0.72, 7.18, 3.2, 0.18, size=8, color=t, bold=True, margin=0)
    add_text(slide, f"{number:02d}", 12.05, 7.16, 0.56, 0.20, size=9, color=t, bold=True,
             align=PP_ALIGN.RIGHT, margin=0)


def add_bullet(slide, text, x, y, w, color=INK, dot=CYAN, size=14, dark=False):
    circle(slide, x, y + 0.12, 0.10, dot)
    add_text(slide, text, x + 0.19, y, w - 0.19, 0.46, size=size, color=color, margin=0, valign=MSO_ANCHOR.TOP)


def add_router(slide, label, x, y, d=0.58, fill=NAVY_2, ring=CYAN, text_color=WHITE):
    circle(slide, x, y, d, fill, ring, 2.0)
    add_text(slide, label, x, y + 0.02, d, d - 0.04, size=12, color=text_color, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_device(slide, label, x, y, w=1.25, color=CYAN, dark=True):
    fill = NAVY_2 if dark else WHITE
    outline = color
    card(slide, x, y, w, 0.66, fill, outline, 1.6)
    # Minimal screen icon
    add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.12, y + 0.13, 0.34, 0.24, color, color)
    line(slide, x + 0.29, y + 0.37, x + 0.29, y + 0.47, color, 1.4)
    line(slide, x + 0.21, y + 0.47, x + 0.37, y + 0.47, color, 1.4)
    add_text(slide, label, x + 0.53, y + 0.09, w - 0.62, 0.45, size=10,
             color=WHITE if dark else INK, bold=True, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_packet(slide, label, x, y, fill=CYAN, text_color=NAVY, w=0.55, h=0.28):
    card(slide, x, y, w, h, fill, fill)
    add_text(slide, label, x, y + 0.01, w, h - 0.02, size=8, color=text_color, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_step_badge(slide, n, x, y, color=CYAN, dark=False):
    circle(slide, x, y, 0.36, color)
    add_text(slide, str(n), x, y + 0.01, 0.36, 0.33, size=11, color=NAVY if dark else WHITE,
             bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # ------------------------------------------------------------------ 1
    slide = prs.slides.add_slide(blank)
    set_bg(slide, NAVY)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.16, H, CYAN, CYAN)
    pill(slide, "COMPUTER NETWORKS • VISUAL GUIDE", 0.78, 0.58, 2.72, CYAN, NAVY, 9)
    add_text(slide, "How Packets\nCross Networks", 0.78, 1.28, 5.9, 1.65, size=38, color=WHITE,
             bold=True, font=FONT_DISPLAY, margin=0, line_spacing=0.90)
    add_text(slide, "Packet topology  •  Datagrams vs. virtual circuits  •  IPv4↔IPv6 translation",
             0.80, 3.20, 5.75, 0.90, size=17, color="B7C4DD", margin=0)
    add_text(slide, "A 10-slide, diagram-first explainer", 0.80, 6.50, 4.2, 0.30, size=11,
             color="70809F", bold=True, margin=0)

    # Hero network visual
    card(slide, 7.08, 0.70, 5.55, 5.95, NAVY_2, "263552", 1.2)
    add_text(slide, "ONE PACKET • THREE QUESTIONS", 7.48, 1.05, 4.75, 0.28, size=10,
             color="8FA0BF", bold=True, margin=0)
    # Links first
    links = [
        (7.88, 2.34, 9.18, 2.03), (7.88, 2.34, 9.20, 3.19),
        (9.48, 2.03, 10.44, 2.60), (9.50, 3.19, 10.44, 2.60),
        (10.75, 2.60, 11.83, 2.29), (10.75, 2.60, 11.80, 3.36),
    ]
    for a, b, c, d in links:
        line(slide, a, b, c, d, "40506E", 2.2)
    # Highlight route
    line(slide, 7.88, 2.34, 9.18, 2.03, CYAN, 3.6)
    line(slide, 9.48, 2.03, 10.44, 2.60, CYAN, 3.6)
    line(slide, 10.75, 2.60, 11.80, 3.36, CYAN, 3.6)
    add_router(slide, "S", 7.58, 2.05, 0.58, NAVY, CYAN)
    add_router(slide, "R1", 9.16, 1.74, 0.58, NAVY, CYAN)
    add_router(slide, "R2", 9.17, 2.91, 0.58, NAVY, "536483")
    add_router(slide, "R3", 10.44, 2.31, 0.58, NAVY, CYAN)
    add_router(slide, "D", 11.78, 3.07, 0.58, NAVY, CYAN)
    add_router(slide, "R4", 11.80, 2.00, 0.58, NAVY, "536483")
    add_packet(slide, "IP", 10.82, 2.72, CYAN, NAVY, 0.48, 0.25)
    # Three question rows
    for idx, (yy, num, label, answer, color) in enumerate([
        (4.12, "01", "TOPOLOGY", "Where can it go?", CYAN),
        (4.78, "02", "SERVICE", "How is the path chosen?", PURPLE),
        (5.44, "03", "TRANSITION", "Can both ends speak the same IP?", MINT),
    ]):
        pill(slide, num, 7.48, yy, 0.46, color, NAVY if color != PURPLE else WHITE, 9)
        add_text(slide, label, 8.15, yy + 0.02, 1.15, 0.20, size=9, color=color, bold=True, margin=0)
        add_text(slide, answer, 9.28, yy - 0.01, 2.75, 0.28, size=11, color=WHITE, bold=True, margin=0)

    # ------------------------------------------------------------------ 2
    slide = prs.slides.add_slide(blank)
    set_bg(slide, PAPER)
    title_block(slide, "Roadmap", "A three-part mental model", "Follow one packet from a device, through a network, to a different IP generation.")
    cards = [
        (0.72, CYAN, CYAN_PALE, "01", "SEE THE NETWORK", "Topology", "Nodes + links reveal every possible route."),
        (4.50, PURPLE, PURPLE_PALE, "02", "CHOOSE THE SERVICE", "Datagram or VC", "The service model controls path behavior."),
        (8.28, MINT, MINT_PALE, "03", "CROSS GENERATIONS", "IPv4 ↔ IPv6", "Transition tools bridge incompatible packet formats."),
    ]
    for x, color, pale, num, kicker, head, body in cards:
        card(slide, x, 2.00, 3.25, 3.82, WHITE, LINE, 1.0)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.00, 3.25, 0.10, color, color)
        circle(slide, x + 0.24, 2.34, 0.56, pale, color, 1.4)
        add_text(slide, num, x + 0.24, 2.37, 0.56, 0.48, size=13, color=color, bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, kicker, x + 0.98, 2.39, 1.95, 0.22, size=9, color=color, bold=True, margin=0)
        add_text(slide, head, x + 0.24, 3.08, 2.72, 0.43, size=22, color=INK, bold=True,
                 font=FONT_DISPLAY, margin=0)
        add_text(slide, body, x + 0.24, 3.65, 2.72, 0.72, size=13, color=SLATE, margin=0)
    # card illustrations
    line(slide, 1.15, 5.03, 1.88, 4.69, "AFC7D1", 2)
    line(slide, 1.88, 4.69, 2.67, 5.10, "AFC7D1", 2)
    line(slide, 1.88, 4.69, 3.22, 4.72, "AFC7D1", 2)
    for x, y in [(1.00, 4.89), (1.73, 4.55), (2.52, 4.96), (3.07, 4.58)]:
        circle(slide, x, y, 0.28, CYAN, CYAN)
    add_packet(slide, "P1", 5.02, 4.70, PURPLE, WHITE, 0.50, 0.26)
    arrow(slide, 5.63, 4.72, 0.50, 0.22, "D0C3F8")
    add_packet(slide, "P2", 6.25, 4.70, PURPLE, WHITE, 0.50, 0.26)
    arrow(slide, 6.86, 4.72, 0.50, 0.22, "D0C3F8")
    pill(slide, "IPv4", 8.77, 4.68, 0.90, AMBER, NAVY, 10)
    arrow(slide, 9.81, 4.72, 0.63, 0.22, MINT)
    pill(slide, "IPv6", 10.58, 4.68, 0.90, MINT, NAVY, 10)
    # arrows between cards
    arrow(slide, 4.08, 3.66, 0.28, 0.42, "CBD4E2")
    arrow(slide, 7.86, 3.66, 0.28, 0.42, "CBD4E2")
    card(slide, 0.72, 6.15, 10.81, 0.60, NAVY, NAVY)
    add_text(slide, "GOAL", 1.00, 6.32, 0.65, 0.22, size=9, color=CYAN, bold=True, margin=0)
    add_text(slide, "Trace one packet end-to-end—and explain every choice along the way.", 1.77, 6.27, 9.25, 0.30,
             size=14, color=WHITE, bold=True, margin=0)
    footer(slide, 2, "Roadmap")

    # ------------------------------------------------------------------ 3
    slide = prs.slides.add_slide(blank)
    set_bg(slide, NAVY)
    title_block(slide, "Part 1 • Topology", "Packet network topology", "Topology is the map: devices are nodes; communication channels are links.", dark=True)
    # Vertical zone labels and backgrounds
    zones = [
        (0.72, 2.00, 2.18, 3.78, "15223D", "EDGE", CYAN),
        (2.98, 2.00, 2.15, 3.78, "171F38", "ACCESS", MINT),
        (5.21, 2.00, 5.10, 3.78, "151B31", "CORE", PURPLE),
        (10.39, 2.00, 2.22, 3.78, "15223D", "EDGE", AMBER),
    ]
    for x, y, w, h, fill, name, accent in zones:
        card(slide, x, y, w, h, fill, "263552", 0.8)
        pill(slide, name, x + 0.18, y + 0.20, 0.76, accent, NAVY, 8)
    # Network links first
    add_device(slide, "Laptop", 1.08, 3.40, 1.45, CYAN, True)
    add_router(slide, "R1", 3.75, 3.40, 0.72, NAVY, MINT)
    core_pts = {"R2": (5.75, 2.64), "R3": (7.30, 3.50), "R4": (8.72, 2.62), "R5": (8.77, 4.52)}
    # Links overlay behind core nodes
    links2 = [
        (2.53, 3.73, 3.75, 3.76), (4.47, 3.76, 5.75, 3.00),
        (4.47, 3.76, 7.30, 3.86), (6.47, 3.00, 7.30, 3.86),
        (6.47, 3.00, 8.72, 2.98), (8.02, 3.86, 8.72, 2.98),
        (8.02, 3.86, 8.77, 4.88), (9.49, 2.98, 10.72, 3.72),
        (9.49, 4.88, 10.72, 3.72),
    ]
    for x1, y1, x2, y2 in links2:
        line(slide, x1, y1, x2, y2, "40506E", 2.4)
    # highlighted selected path
    for x1, y1, x2, y2 in [(2.53, 3.73, 3.75, 3.76), (4.47, 3.76, 5.75, 3.00),
                            (6.47, 3.00, 8.72, 2.98), (9.49, 2.98, 10.72, 3.72)]:
        line(slide, x1, y1, x2, y2, CYAN, 3.8)
    for label, (x, y) in core_pts.items():
        add_router(slide, label, x, y, 0.72, NAVY, CYAN if label in ("R2", "R4") else "536483")
    add_device(slide, "Server", 10.72, 3.39, 1.50, AMBER, True)
    add_packet(slide, "P", 7.70, 2.78, CYAN, NAVY, 0.42, 0.24)
    # glossary strip
    terms = [
        (0.92, "NODE", "A host, switch, or router", CYAN),
        (3.45, "LINK", "A wired or wireless channel", MINT),
        (6.25, "PATH", "A sequence of links", PURPLE),
        (8.92, "HOP", "One router-to-router step", AMBER),
    ]
    for x, head, desc, color in terms:
        add_text(slide, head, x, 6.05, 0.58, 0.22, size=9, color=color, bold=True, margin=0)
        add_text(slide, desc, x + 0.67, 6.02, 1.76, 0.34, size=10, color="B7C4DD", margin=0)
    card(slide, 0.72, 6.56, 11.89, 0.41, "18243E", "263552")
    add_text(slide, "MAP ≠ DECISION", 0.93, 6.67, 1.18, 0.18, size=8, color=CYAN, bold=True, margin=0)
    add_text(slide, "Topology lists possible paths; routing chooses one path for this packet.", 2.20, 6.61, 9.70, 0.26,
             size=11, color=WHITE, bold=True, margin=0)
    footer(slide, 3, "Packet network topology", dark=True)

    # ------------------------------------------------------------------ 4
    slide = prs.slides.add_slide(blank)
    set_bg(slide, PAPER)
    title_block(slide, "Part 1 • Forwarding", "What happens at each hop?", "Data is wrapped for delivery, then each router repeats a small decision loop.")
    stages = [
        (0.72, 2.02, 2.30, CYAN, CYAN_PALE, "1", "APPLICATION DATA", "Hello"),
        (3.28, 2.02, 2.30, PURPLE, PURPLE_PALE, "2", "TRANSPORT SEGMENT", "Ports  |  Data"),
        (5.84, 2.02, 2.30, MINT, MINT_PALE, "3", "IP DATAGRAM", "Src/Dst IP  |  Payload"),
        (8.40, 2.02, 2.30, AMBER, AMBER_PALE, "4", "LINK FRAME", "MAC  |  IP packet  |  FCS"),
    ]
    for x, y, w, color, pale, num, head, sample in stages:
        card(slide, x, y, w, 1.62, WHITE, LINE)
        add_step_badge(slide, num, x + 0.18, y + 0.18, color)
        add_text(slide, head, x + 0.68, y + 0.22, 1.43, 0.23, size=9, color=color, bold=True, margin=0)
        card(slide, x + 0.18, y + 0.78, w - 0.36, 0.47, pale, pale)
        add_text(slide, sample, x + 0.24, y + 0.88, w - 0.48, 0.24, size=11, color=INK, bold=True,
                 align=PP_ALIGN.CENTER, margin=0)
    for x in (3.02, 5.58, 8.14):
        arrow(slide, x, 2.62, 0.20, 0.32, "C9D2E0")

    add_text(slide, "ROUTER LOOP", 0.72, 4.05, 1.10, 0.24, size=9, color=CYAN_DARK, bold=True, margin=0)
    card(slide, 0.72, 4.42, 11.89, 1.38, NAVY, NAVY)
    processes = [
        (1.02, "1", "READ", "destination IP", CYAN),
        (3.37, "2", "LOOK UP", "forwarding table", PURPLE),
        (5.86, "3", "CHOOSE", "next hop", MINT),
        (8.13, "4", "UPDATE", "TTL / Hop Limit", AMBER),
        (10.38, "5", "RE-FRAME", "for the next link", CORAL),
    ]
    for i, (x, num, head, sub, color) in enumerate(processes):
        circle(slide, x, 4.77, 0.48, color)
        add_text(slide, num, x, 4.83, 0.48, 0.28, size=10, color=NAVY if color != PURPLE else WHITE,
                 bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, head, x + 0.61, 4.66, 1.15, 0.25, size=10, color=color, bold=True, margin=0)
        add_text(slide, sub, x + 0.61, 4.95, 1.30, 0.34, size=10, color=WHITE, margin=0)
        if i < len(processes) - 1:
            arrow(slide, x + 1.86, 4.91, 0.26, 0.24, "3A4863")
    card(slide, 0.72, 6.10, 11.89, 0.66, WHITE, LINE)
    pill(slide, "KEY IDEA", 0.94, 6.25, 0.92, CYAN, NAVY, 8)
    add_text(slide, "IP addresses normally stay end-to-end; every router decreases TTL/Hop Limit and builds a new link frame.",
             2.04, 6.19, 9.95, 0.36, size=12, color=INK, bold=True, margin=0)
    footer(slide, 4, "Hop-by-hop forwarding")

    # ------------------------------------------------------------------ 5
    slide = prs.slides.add_slide(blank)
    set_bg(slide, NAVY)
    title_block(slide, "Part 2 • Service model", "Datagram network: every packet decides", "No connection setup. Each packet carries a destination and can be routed independently.", dark=True)
    # Diagram panel
    card(slide, 0.72, 1.92, 8.10, 4.74, NAVY_2, "263552")
    add_device(slide, "Source", 1.10, 3.60, 1.40, CYAN, True)
    add_device(slide, "Destination", 7.02, 3.60, 1.46, MINT, True)
    pts = {"R1": (3.00, 2.45), "R2": (3.00, 4.75), "R3": (4.62, 3.55), "R4": (6.05, 2.45), "R5": (6.05, 4.75)}
    all_links = [
        (2.50, 3.93, 3.00, 2.79), (2.50, 3.93, 3.00, 5.09),
        (3.62, 2.79, 4.62, 3.89), (3.62, 5.09, 4.62, 3.89),
        (5.24, 3.89, 6.05, 2.79), (5.24, 3.89, 6.05, 5.09),
        (6.67, 2.79, 7.02, 3.93), (6.67, 5.09, 7.02, 3.93),
    ]
    for a, b, c, d in all_links:
        line(slide, a, b, c, d, "384761", 2.3)
    # Two route overlays
    for a, b, c, d in [(2.50, 3.84, 3.00, 2.70), (3.62, 2.70, 4.62, 3.80),
                        (5.24, 3.80, 6.05, 2.70), (6.67, 2.70, 7.02, 3.84)]:
        line(slide, a, b, c, d, CYAN, 3.5)
    for a, b, c, d in [(2.50, 4.03, 3.00, 5.18), (3.62, 5.18, 4.62, 3.98),
                        (5.24, 3.98, 6.05, 5.18), (6.67, 5.18, 7.02, 4.03)]:
        line(slide, a, b, c, d, AMBER, 3.5)
    for name, (x, y) in pts.items():
        add_router(slide, name, x, y, 0.62, NAVY, "71809A")
    add_packet(slide, "P1", 4.00, 2.72, CYAN, NAVY, 0.46, 0.25)
    add_packet(slide, "P2", 4.00, 4.82, AMBER, NAVY, 0.46, 0.25)
    add_text(slide, "PACKET 1 • upper path", 1.05, 2.17, 1.70, 0.22, size=9, color=CYAN, bold=True, margin=0)
    add_text(slide, "PACKET 2 • lower path", 6.31, 5.95, 1.82, 0.22, size=9, color=AMBER, bold=True, margin=0)
    # Explanation panel
    card(slide, 9.10, 1.92, 3.51, 4.74, "151F37", "263552")
    pill(slide, "CONNECTIONLESS", 9.43, 2.24, 1.42, CYAN, NAVY, 8)
    bullets = [
        ("Send immediately—no path setup", CYAN),
        ("Full destination in each packet", PURPLE),
        ("Routers use forwarding tables", MINT),
        ("Packets may take different paths", AMBER),
        ("Loss or reordering is possible", CORAL),
    ]
    yy = 2.91
    for text_value, color in bullets:
        add_bullet(slide, text_value, 9.45, yy, 2.78, color=WHITE, dot=color, size=12)
        yy += 0.62
    card(slide, 9.42, 6.04, 2.90, 0.35, "20304F", "20304F")
    add_text(slide, "The Internet Protocol (IP) uses datagrams.", 9.55, 6.10, 2.62, 0.20, size=9,
             color="B7C4DD", bold=True, align=PP_ALIGN.CENTER, margin=0)
    footer(slide, 5, "Datagram networks", dark=True)

    # ------------------------------------------------------------------ 6
    slide = prs.slides.add_slide(blank)
    set_bg(slide, PAPER)
    title_block(slide, "Part 2 • Service model", "Virtual circuit: set the path first", "A logical path is established once; later packets carry a short label and follow that path.")
    # Lifecycle
    steps = [
        (0.72, "1", "SETUP", "Choose path + install state", CYAN, CYAN_PALE),
        (4.50, "2", "TRANSFER", "Forward using VC labels", PURPLE, PURPLE_PALE),
        (8.28, "3", "TEARDOWN", "Release circuit state", MINT, MINT_PALE),
    ]
    for x, num, head, sub, color, pale in steps:
        card(slide, x, 1.90, 3.25, 0.88, WHITE, LINE)
        circle(slide, x + 0.18, 2.14, 0.40, color)
        add_text(slide, num, x + 0.18, 2.18, 0.40, 0.27, size=10, color=NAVY if color != PURPLE else WHITE,
                 bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, head, x + 0.72, 2.08, 1.02, 0.22, size=10, color=color, bold=True, margin=0)
        add_text(slide, sub, x + 0.72, 2.34, 2.18, 0.24, size=10, color=SLATE, margin=0)
    arrow(slide, 4.07, 2.17, 0.28, 0.28, "CBD4E2")
    arrow(slide, 7.85, 2.17, 0.28, 0.28, "CBD4E2")

    # Fixed path diagram
    card(slide, 0.72, 3.15, 8.15, 3.45, WHITE, LINE)
    add_text(slide, "FIXED LOGICAL PATH", 1.00, 3.42, 1.72, 0.23, size=9, color=PURPLE, bold=True, margin=0)
    positions = [(1.05, "S", CYAN), (2.55, "R1", PURPLE), (4.15, "R2", PURPLE), (5.75, "R3", PURPLE), (7.35, "D", MINT)]
    for i in range(len(positions) - 1):
        x1 = positions[i][0] + 0.66
        x2 = positions[i + 1][0]
        line(slide, x1, 4.52, x2, 4.52, PURPLE, 4.0)
    labels = ["VC 42", "VC 17", "VC 8", "VC 31"]
    for i, (x, label, color) in enumerate(positions):
        add_router(slide, label, x, 4.19, 0.66, WHITE, color, INK)
        if i < len(labels):
            pill(slide, labels[i], x + 0.78, 4.12, 0.62, PURPLE_PALE, PURPLE_DARK, 7, PURPLE_PALE)
    # unused alternative path
    line(slide, 2.88, 4.85, 4.30, 5.70, "D7DDEA", 1.7, MSO_LINE_DASH_STYLE.DASH)
    line(slide, 4.62, 5.70, 6.08, 4.85, "D7DDEA", 1.7, MSO_LINE_DASH_STYLE.DASH)
    add_router(slide, "R4", 4.29, 5.41, 0.60, WHITE, "B8C1D0", MUTED)
    add_text(slide, "other route not used", 3.80, 6.12, 1.60, 0.20, size=8, color=MUTED,
             align=PP_ALIGN.CENTER, margin=0)
    # Properties
    card(slide, 9.12, 3.15, 3.49, 3.45, NAVY, NAVY)
    pill(slide, "PATH-ORIENTED", 9.43, 3.48, 1.31, PURPLE, WHITE, 8)
    props = [
        "Switches keep circuit state",
        "Packets follow one chosen path",
        "Delivery is typically ordered",
        "A failure may trigger re-setup",
    ]
    yy = 4.12
    for p in props:
        add_bullet(slide, p, 9.45, yy, 2.75, color=WHITE, dot=PURPLE, size=11)
        yy += 0.52
    add_text(slide, "Examples", 9.44, 6.05, 0.62, 0.20, size=9, color=PURPLE, bold=True, margin=0)
    add_text(slide, "ATM / Frame Relay • MPLS is VC-like", 10.10, 6.02, 2.05, 0.30, size=9, color="B7C4DD", margin=0)
    footer(slide, 6, "Virtual circuits")

    # ------------------------------------------------------------------ 7
    slide = prs.slides.add_slide(blank)
    set_bg(slide, NAVY)
    title_block(slide, "Part 2 • Compare", "Datagrams vs. virtual circuits", "Both move packets. The difference is when—and where—the path decision is remembered.", dark=True)
    # Column headers
    card(slide, 0.72, 1.88, 3.04, 0.82, "16223C", "263552")
    add_text(slide, "QUESTION", 1.00, 2.13, 2.44, 0.25, size=10, color="8FA0BF", bold=True, margin=0)
    card(slide, 3.93, 1.88, 4.16, 0.82, "112C43", CYAN_DARK)
    circle(slide, 4.20, 2.10, 0.35, CYAN)
    add_text(slide, "D", 4.20, 2.14, 0.35, 0.22, size=9, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "DATAGRAM", 4.72, 2.09, 1.40, 0.28, size=13, color=CYAN, bold=True, margin=0)
    add_text(slide, "decide per packet", 6.33, 2.11, 1.38, 0.24, size=10, color="B7C4DD", margin=0)
    card(slide, 8.26, 1.88, 4.35, 0.82, "251D48", PURPLE_DARK)
    circle(slide, 8.54, 2.10, 0.35, PURPLE)
    add_text(slide, "V", 8.54, 2.14, 0.35, 0.22, size=9, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "VIRTUAL CIRCUIT", 9.05, 2.09, 1.73, 0.28, size=13, color="B9A0FF", bold=True, margin=0)
    add_text(slide, "decide before data", 10.88, 2.11, 1.42, 0.24, size=10, color="B7C4DD", margin=0)

    rows = [
        ("Before sending", "Nothing to set up", "Signaling creates a path"),
        ("Packet carries", "Full destination address", "Short VC / label"),
        ("Network remembers", "Forwarding table", "Forwarding + circuit state"),
        ("Path & order", "Path may vary; reorder possible", "Usually fixed; usually ordered"),
        ("After a failure", "Later packets can reroute", "Circuit may need rebuilding"),
        ("Best fit", "Internet-scale flexibility", "Managed paths / traffic engineering"),
    ]
    y = 2.84
    for i, (question, dgram, vc) in enumerate(rows):
        fill = "121B30" if i % 2 == 0 else "0F172B"
        card(slide, 0.72, y, 11.89, 0.55, fill, fill)
        add_text(slide, question, 1.00, y + 0.14, 2.55, 0.22, size=10, color="91A0BD", bold=True, margin=0)
        add_text(slide, dgram, 4.19, y + 0.13, 3.60, 0.24, size=11, color=WHITE, margin=0)
        add_text(slide, vc, 8.52, y + 0.13, 3.70, 0.24, size=11, color=WHITE, margin=0)
        y += 0.58

    # Bottom decision cards
    card(slide, 0.72, 6.48, 5.75, 0.48, "11304A", CYAN_DARK)
    add_text(slide, "Choose DATAGRAM for resilience + massive scale", 0.98, 6.60, 5.20, 0.23,
             size=11, color=CYAN, bold=True, align=PP_ALIGN.CENTER, margin=0)
    card(slide, 6.63, 6.48, 5.98, 0.48, "271F4B", PURPLE_DARK)
    add_text(slide, "Choose VC-LIKE paths for control + predictability", 6.90, 6.60, 5.40, 0.23,
             size=11, color="C7B5FF", bold=True, align=PP_ALIGN.CENTER, margin=0)
    footer(slide, 7, "Datagram vs. virtual circuit", dark=True)

    # ------------------------------------------------------------------ 8
    slide = prs.slides.add_slide(blank)
    set_bg(slide, PAPER)
    title_block(slide, "Part 3 • Transition", "Why IPv4 and IPv6 need a bridge", "They solve the same job, but use different address sizes and network-layer headers.")
    # Side cards
    card(slide, 0.72, 1.90, 4.46, 3.45, WHITE, LINE)
    pill(slide, "IPv4", 1.03, 2.20, 0.86, AMBER, NAVY, 11)
    add_text(slide, "32-bit addresses", 1.03, 2.80, 3.54, 0.36, size=24, color=INK, bold=True,
             font=FONT_DISPLAY, margin=0)
    card(slide, 1.03, 3.36, 3.84, 0.58, AMBER_PALE, AMBER_PALE)
    add_text(slide, "192.0.2.34", 1.19, 3.48, 3.50, 0.30, size=18, color=AMBER_DARK, bold=True,
             font="Consolas", align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "≈ 4.3 billion unique addresses", 1.03, 4.23, 3.60, 0.28, size=12, color=SLATE, margin=0)
    add_text(slide, "Shortage drove widespread NAT use.", 1.03, 4.62, 3.62, 0.30, size=11, color=MUTED, margin=0)

    card(slide, 8.15, 1.90, 4.46, 3.45, WHITE, LINE)
    pill(slide, "IPv6", 8.46, 2.20, 0.86, MINT, NAVY, 11)
    add_text(slide, "128-bit addresses", 8.46, 2.80, 3.54, 0.36, size=24, color=INK, bold=True,
             font=FONT_DISPLAY, margin=0)
    card(slide, 8.46, 3.36, 3.84, 0.58, MINT_PALE, MINT_PALE)
    add_text(slide, "2001:db8::34", 8.62, 3.48, 3.50, 0.30, size=17, color=MINT_DARK, bold=True,
             font="Consolas", align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "≈ 3.4 × 10³⁸ unique addresses", 8.46, 4.23, 3.60, 0.28, size=12, color=SLATE, margin=0)
    add_text(slide, "Enough space for end-to-end addressing.", 8.46, 4.62, 3.62, 0.30, size=11, color=MUTED, margin=0)

    # Middle bridge / incompatibility
    circle(slide, 5.76, 2.45, 1.82, NAVY)
    add_text(slide, "≠", 5.76, 2.63, 1.82, 0.85, size=38, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, "different packet\nformats", 5.91, 3.45, 1.52, 0.60, size=10, color=SLATE, bold=True,
             align=PP_ALIGN.CENTER, margin=0)

    add_text(slide, "THREE WAYS TO CROSS", 0.72, 5.73, 2.06, 0.22, size=9, color=CYAN_DARK, bold=True, margin=0)
    transition_cards = [
        (0.72, CYAN, CYAN_PALE, "DUAL STACK", "Run IPv4 + IPv6 together", "both"),
        (4.72, PURPLE, PURPLE_PALE, "TUNNEL", "Encapsulate one inside the other", "wrap"),
        (8.72, MINT, MINT_PALE, "TRANSLATE", "Rewrite between packet formats", "map"),
    ]
    for x, color, pale, head, body, tag in transition_cards:
        card(slide, x, 6.08, 3.55, 0.76, WHITE, LINE)
        pill(slide, tag.upper(), x + 0.18, 6.29, 0.62, color, NAVY if color != PURPLE else WHITE, 7)
        add_text(slide, head, x + 0.95, 6.17, 1.12, 0.22, size=10, color=color, bold=True, margin=0)
        add_text(slide, body, x + 0.95, 6.42, 2.34, 0.22, size=9, color=SLATE, margin=0)
    footer(slide, 8, "IPv4 and IPv6")

    # ------------------------------------------------------------------ 9
    slide = prs.slides.add_slide(blank)
    set_bg(slide, NAVY)
    title_block(slide, "Part 3 • Translation", "Translation toolbox: match the constraint", "Translation is not one technique—state, direction, and application needs determine the right choice.", dark=True)
    cards9 = [
        (0.72, CYAN, "STATELESS • SIIT", "Algorithmic map", "Converts IPv4 ↔ IPv6 headers and addresses without a session table.",
         "GOOD FOR", "Predictable 1:1 or prefix-based mappings"),
        (4.70, PURPLE, "STATEFUL • NAT64 + DNS64", "Share IPv4 reachability", "DNS64 creates an IPv6 destination; NAT64 stores sessions and translates packets.",
         "GOOD FOR", "IPv6-only clients reaching IPv4 servers"),
        (8.68, MINT, "DOUBLE TRANSLATION • 464XLAT", "Keep legacy apps working", "CLAT translates IPv4→IPv6 near the client; PLAT/NAT64 translates IPv6→IPv4.",
         "GOOD FOR", "IPv4-only apps on IPv6-only access networks"),
    ]
    for x, color, label, head, desc, good, fit in cards9:
        card(slide, x, 1.92, 3.65, 4.52, "151F37", "263552")
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 1.92, 3.65, 0.09, color, color)
        pill(slide, label, x + 0.25, 2.26, 2.24, color, NAVY if color != PURPLE else WHITE, 8)
        add_text(slide, head, x + 0.25, 2.91, 3.12, 0.66, size=20, color=WHITE, bold=True,
                 font=FONT_DISPLAY, margin=0)
        add_text(slide, desc, x + 0.25, 3.70, 3.10, 1.04, size=12, color="B7C4DD", margin=0)
        # tiny flow visual
        pill(slide, "v4", x + 0.28, 4.91, 0.51, AMBER, NAVY, 8)
        arrow(slide, x + 0.94, 4.97, 0.40, 0.20, color)
        pill(slide, "T", x + 1.49, 4.91, 0.51, color, NAVY if color != PURPLE else WHITE, 8)
        arrow(slide, x + 2.15, 4.97, 0.40, 0.20, color)
        pill(slide, "v6", x + 2.70, 4.91, 0.51, MINT, NAVY, 8)
        line(slide, x + 0.25, 5.55, x + 3.39, 5.55, "31405C", 0.8)
        add_text(slide, good, x + 0.25, 5.75, 0.74, 0.20, size=8, color=color, bold=True, margin=0)
        add_text(slide, fit, x + 1.06, 5.69, 2.22, 0.48, size=10, color=WHITE, bold=True, margin=0)
    card(slide, 0.72, 6.67, 11.61, 0.30, "1B2945", "263552")
    add_text(slide, "WATCH FOR", 0.93, 6.73, 0.88, 0.16, size=8, color=CORAL, bold=True, margin=0)
    add_text(slide, "IP literals • embedded addresses • IPsec behavior • logging and troubleshooting across mapped addresses",
             1.92, 6.69, 10.14, 0.20, size=9, color="B7C4DD", margin=0)
    footer(slide, 9, "IPv4–IPv6 translation strategies", dark=True)

    # ----------------------------------------------------------------- 10
    slide = prs.slides.add_slide(blank)
    set_bg(slide, PAPER)
    title_block(slide, "Part 3 • Worked flow", "NAT64 + DNS64: one packet, five moves", "Example: an IPv6-only client opens a website hosted on an IPv4-only server.")
    # Main flow lane
    card(slide, 0.72, 1.88, 11.89, 3.16, WHITE, LINE)
    nodes = [
        (1.00, 2.64, 1.62, CYAN, "IPv6 client", "2001:db8::10"),
        (3.26, 2.64, 1.62, PURPLE, "DNS64", "synthesizes AAAA"),
        (5.52, 2.64, 1.62, AMBER, "IPv6 target", "64:ff9b::c000:222"),
        (7.78, 2.64, 1.62, MINT, "NAT64", "keeps session state"),
        (10.04, 2.64, 1.62, AMBER, "IPv4 server", "192.0.2.34"),
    ]
    for idx in range(len(nodes) - 1):
        arrow(slide, nodes[idx][0] + 1.78, 3.13, 0.34, 0.30, "C9D2E0")
    for idx, (x, y, w, color, head, sub) in enumerate(nodes, start=1):
        card(slide, x, y, w, 1.34, "FBFCFE", LINE)
        add_step_badge(slide, idx, x + 0.12, y + 0.12, color)
        add_text(slide, head, x + 0.12, y + 0.58, w - 0.24, 0.26, size=11, color=INK, bold=True,
                 align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, sub, x + 0.10, y + 0.90, w - 0.20, 0.28, size=8 if idx == 3 else 9,
                 color=SLATE, font="Consolas" if idx in (1, 3, 5) else FONT,
                 align=PP_ALIGN.CENTER, margin=0)
    # move descriptions below lane
    descriptions = [
        (0.95, "1", "Ask for example.com"),
        (3.21, "2", "Build AAAA from A record"),
        (5.47, "3", "Send an IPv6 packet"),
        (7.73, "4", "Translate v6 header → v4"),
        (9.99, "5", "Reply is translated back"),
    ]
    for x, n, desc in descriptions:
        add_text(slide, desc, x, 4.33, 1.78, 0.34, size=9, color=MUTED, bold=True,
                 align=PP_ALIGN.CENTER, margin=0)

    # Decision strip
    add_text(slide, "QUICK STRATEGY CHOOSER", 0.72, 5.45, 2.45, 0.22, size=9, color=CYAN_DARK, bold=True, margin=0)
    chooser = [
        (0.72, 5.82, 2.67, CYAN, "Both ends support IPv6?", "Native IPv6 / dual stack"),
        (3.52, 5.82, 2.67, PURPLE, "IPv6 client → IPv4 server?", "NAT64 + DNS64"),
        (6.32, 5.82, 2.67, MINT, "IPv4-only app on IPv6 access?", "464XLAT"),
        (9.12, 5.82, 3.49, AMBER, "IPv6 islands over an IPv4 core?", "Tunnel"),
    ]
    for x, y, w, color, q, a in chooser:
        card(slide, x, y, w, 0.84, WHITE, LINE)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.08, 0.84, color, color)
        add_text(slide, q, x + 0.23, y + 0.13, w - 0.37, 0.22, size=9, color=SLATE, bold=True, margin=0)
        add_text(slide, a, x + 0.23, y + 0.46, w - 0.37, 0.22, size=11, color=color if color != AMBER else AMBER_DARK,
                 bold=True, margin=0)
    add_text(slide, "Remember: topology shows the choices • the service model shapes forwarding • translation connects IP generations",
             0.72, 6.84, 11.89, 0.24, size=10, color=INK, bold=True, align=PP_ALIGN.CENTER, margin=0)
    footer(slide, 10, "NAT64 + DNS64 worked example")

    # Document metadata
    prs.core_properties.title = "Packet Networks and IPv4-to-IPv6 Translation"
    prs.core_properties.subject = "Packet topology, datagrams, virtual circuits, and IPv4/IPv6 translation strategies"
    prs.core_properties.author = "OpenAI"
    prs.core_properties.keywords = "computer networks, IPv4, IPv6, NAT64, DNS64, datagram, virtual circuit"
    prs.core_properties.comments = "Diagram-first educational presentation; 10 slides."
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    output = build_deck()
    print(output)
