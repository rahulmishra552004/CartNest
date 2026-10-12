"""Build a phone-safe, 12-slide networking presentation with white backgrounds.

The slides are drawn at 2560x1440 and embedded as full-slide PNGs. This keeps
the visual layout consistent in mobile PowerPoint viewers.
"""

from __future__ import annotations

import math
from functools import lru_cache
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


WIDTH, HEIGHT = 2560, 1440
OUT = Path("/tmp/computer_networks_complete_visual_guide.pptx")
ASSET_DIR = Path("/tmp/computer_networks_white_slides")

WHITE = "#FFFFFF"
INK = "#172033"
GRAY = "#667085"
LIGHT = "#E8ECF2"
PALE = "#F6F8FB"
BLUE = "#2855D9"
BLUE_PALE = "#E9EFFF"
CORAL = "#F05B4F"
CORAL_PALE = "#FDEAE7"
GREEN = "#159A78"
GREEN_PALE = "#E5F6F1"
YELLOW = "#F3C63F"
YELLOW_PALE = "#FFF5D6"
PURPLE = "#7A55C7"
PURPLE_PALE = "#F0EBFA"

FONT_REG = Path("/mnt/c/Windows/Fonts/segoeui.ttf")
FONT_BOLD = Path("/mnt/c/Windows/Fonts/segoeuib.ttf")
FONT_SERIF = Path("/mnt/c/Windows/Fonts/georgia.ttf")
FONT_SERIF_BOLD = Path("/mnt/c/Windows/Fonts/georgiab.ttf")
FONT_MONO = Path("/mnt/c/Windows/Fonts/consola.ttf")
FONT_MONO_BOLD = Path("/mnt/c/Windows/Fonts/consolab.ttf")


@lru_cache(maxsize=None)
def get_font(size: int, family: str = "sans", bold: bool = False):
    if family == "serif":
        path = FONT_SERIF_BOLD if bold else FONT_SERIF
    elif family == "mono":
        path = FONT_MONO_BOLD if bold else FONT_MONO
    else:
        path = FONT_BOLD if bold else FONT_REG
    return ImageFont.truetype(str(path), size)


def text_width(draw, text, font):
    return draw.textbbox((0, 0), text, font=font)[2]


def wrap_lines(draw, text, font, max_width):
    lines = []
    for hard_line in str(text).split("\n"):
        words = hard_line.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            candidate = f"{current} {word}"
            if text_width(draw, candidate, font) <= max_width:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def draw_text(draw, text, box, size=30, color=INK, bold=False, family="sans",
              align="left", valign="top", spacing=8):
    x1, y1, x2, y2 = box
    font = get_font(size, family, bold)
    lines = wrap_lines(draw, text, font, x2 - x1)
    metrics = []
    for line_value in lines:
        bbox = draw.textbbox((0, 0), line_value or "Ag", font=font)
        metrics.append((line_value, bbox[2] - bbox[0], bbox[3] - bbox[1]))
    total_h = sum(h for _, _, h in metrics) + spacing * max(0, len(metrics) - 1)
    if valign == "middle":
        y = y1 + (y2 - y1 - total_h) / 2
    elif valign == "bottom":
        y = y2 - total_h
    else:
        y = y1
    for line_value, line_w, line_h in metrics:
        if align == "center":
            x = x1 + (x2 - x1 - line_w) / 2
        elif align == "right":
            x = x2 - line_w
        else:
            x = x1
        draw.text((x, y), line_value, font=font, fill=color)
        y += line_h + spacing
    return y


def rule(draw, x1, y1, x2, y2, color=LIGHT, width=3):
    draw.line((x1, y1, x2, y2), fill=color, width=width)


def arrow(draw, start, end, color=INK, width=6, head=18, dashed=False):
    x1, y1 = start
    x2, y2 = end
    if dashed:
        length = math.hypot(x2 - x1, y2 - y1)
        ux, uy = (x2 - x1) / length, (y2 - y1) / length
        pos = 0
        while pos < length - head:
            seg_end = min(pos + 22, length - head)
            draw.line((x1 + ux * pos, y1 + uy * pos, x1 + ux * seg_end, y1 + uy * seg_end),
                      fill=color, width=width)
            pos += 38
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    left = (x2 - head * math.cos(angle - math.pi / 6), y2 - head * math.sin(angle - math.pi / 6))
    right = (x2 - head * math.cos(angle + math.pi / 6), y2 - head * math.sin(angle + math.pi / 6))
    draw.polygon([(x2, y2), left, right], fill=color)


def panel(draw, box, fill=WHITE, outline=INK, width=3, shadow=False):
    x1, y1, x2, y2 = box
    if shadow:
        draw.rectangle((x1 + 12, y1 + 12, x2 + 12, y2 + 12), fill="#E9EDF3")
    draw.rectangle(box, fill=fill, outline=outline, width=width)


def tag(draw, text, box, fill=BLUE, color=WHITE, size=22):
    draw.rectangle(box, fill=fill)
    draw_text(draw, text.upper(), box, size=size, color=color, bold=True,
              align="center", valign="middle", spacing=0)


def numbered_badge(draw, n, center, fill=BLUE, diameter=52, color=WHITE):
    x, y = center
    r = diameter / 2
    draw.ellipse((x - r, y - r, x + r, y + r), fill=fill)
    draw_text(draw, str(n), (x - r, y - r, x + r, y + r), size=23, color=color,
              bold=True, align="center", valign="middle", spacing=0)


def packet(draw, box, label, fill=BLUE, color=WHITE):
    draw.rectangle(box, fill=fill, outline=INK, width=2)
    draw_text(draw, label, box, size=19, color=color, bold=True,
              align="center", valign="middle", spacing=0)


def router(draw, center, label, accent=BLUE, radius=46):
    x, y = center
    draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=WHITE, outline=accent, width=6)
    draw_text(draw, label, (x - radius, y - radius, x + radius, y + radius), size=23, color=INK,
              bold=True, align="center", valign="middle", spacing=0)


def switch(draw, box, label="SW", accent=GREEN):
    x1, y1, x2, y2 = box
    draw.rectangle(box, fill=WHITE, outline=accent, width=6)
    for i in range(5):
        x = x1 + 20 + i * ((x2 - x1 - 40) / 4)
        draw.ellipse((x - 5, y2 - 17, x + 5, y2 - 7), fill=accent)
    draw_text(draw, label, (x1, y1 + 2, x2, y2 - 17), size=20, color=INK, bold=True,
              align="center", valign="middle", spacing=0)


def base_slide(number, section, title, subtitle, accent=BLUE):
    image = Image.new("RGB", (WIDTH, HEIGHT), WHITE)
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 34, HEIGHT), fill=accent)
    draw_text(draw, section.upper(), (105, 62, 1650, 105), size=22, color=accent, bold=True)
    title_size = 62 if len(title) > 38 else 68
    draw_text(draw, title, (105, 112, 2140, 205), size=title_size, color=INK, bold=True, family="serif")
    draw_text(draw, subtitle, (110, 220, 2200, 270), size=27, color=GRAY)
    draw_text(draw, f"{number:02d}", (2220, 50, 2470, 170), size=86, color="#D9DEE7",
              bold=True, family="serif", align="right")
    rule(draw, 105, 292, 2460, 292, color=INK, width=3)
    rule(draw, 105, 1378, 2460, 1378, color=LIGHT, width=2)
    draw_text(draw, "COMPUTER NETWORKS • VISUAL NOTES", (105, 1392, 900, 1425), size=17, color=GRAY, bold=True)
    draw_text(draw, f"{number} / 12", (2200, 1392, 2460, 1425), size=17, color=GRAY, bold=True, align="right")
    return image, draw


def slide_01():
    image = Image.new("RGB", (WIDTH, HEIGHT), WHITE)
    d = ImageDraw.Draw(image)
    d.rectangle((0, 0, 44, HEIGHT), fill=CORAL)
    tag(d, "12-SLIDE VISUAL GUIDE", (130, 95, 560, 145), fill=BLUE, size=20)
    draw_text(d, "FROM PREFIXES\nTO PACKETS", (130, 230, 1240, 580), size=102, color=INK,
              bold=True, family="serif", spacing=8)
    draw_text(d, "CIDR • NAT • IPv4↔IPv6 • Topology • Datagram & VC • Switch/Router",
              (138, 650, 1210, 745), size=31, color=GRAY, spacing=6)
    draw_text(d, "A white-background, diagram-first guide to how modern networks plan, forward, and translate traffic.",
              (138, 800, 1110, 930), size=31, color=INK, bold=True, spacing=8)
    rule(d, 138, 1015, 1110, 1015, color=INK, width=3)
    draw_text(d, "COMPUTER NETWORKS", (138, 1050, 780, 1095), size=22, color=CORAL, bold=True)

    # A completely different editorial "metro map" visual.
    d.ellipse((1380, 150, 2400, 1170), outline=LIGHT, width=5)
    path = [(1450, 340), (1740, 340), (1740, 580), (2040, 580), (2040, 820), (2310, 820)]
    colors = [BLUE, GREEN, CORAL, PURPLE, YELLOW]
    for i in range(len(path) - 1):
        arrow(d, path[i], path[i + 1], color=colors[i], width=10, head=26)
    stops = [
        (1450, 340, "/26", "CIDR", BLUE),
        (1740, 340, "N", "NAT", GREEN),
        (1740, 580, "6", "IPv6", CORAL),
        (2040, 580, "T", "Topology", PURPLE),
        (2040, 820, "D", "Datagram / VC", YELLOW),
        (2310, 820, "R", "Switch / Router", BLUE),
    ]
    for x, y, code, label_value, color in stops:
        d.ellipse((x - 47, y - 47, x + 47, y + 47), fill=WHITE, outline=color, width=9)
        draw_text(d, code, (x - 45, y - 45, x + 45, y + 45), size=23, color=INK, bold=True,
                  align="center", valign="middle", spacing=0)
        draw_text(d, label_value, (x - 90, y + 65, x + 130, y + 110), size=20, color=INK, bold=True,
                  align="center")
    draw_text(d, "PLAN", (1410, 1190, 1580, 1230), size=19, color=BLUE, bold=True)
    arrow(d, (1600, 1205), (2240, 1205), color=INK, width=4, head=18)
    draw_text(d, "DELIVER", (2260, 1190, 2430, 1230), size=19, color=CORAL, bold=True)
    draw_text(d, "01", (2260, 1280, 2440, 1380), size=76, color="#D9DEE7", bold=True, family="serif", align="right")
    return image


def slide_02():
    image, d = base_slide(2, "Orientation", "One packet, six questions", "Use this sequence to connect every topic in the deck.", BLUE)
    items = [
        ("01", "PLAN", "Which prefix fits?", "CIDR defines network and host bits.", BLUE),
        ("02", "DIVIDE", "How should space be split?", "VLSM assigns right-sized subnets.", GREEN),
        ("03", "TRANSLATE", "Private or public?", "NAT rewrites addresses and ports.", CORAL),
        ("04", "BRIDGE", "IPv4 or IPv6?", "Transition tools connect generations.", PURPLE),
        ("05", "MOVE", "Which path and service?", "Topology plus datagram/VC behavior.", YELLOW),
        ("06", "FORWARD", "What happens inside?", "Ports, lookup, fabric, queues.", BLUE),
    ]
    xs = [115, 910, 1705]
    ys = [360, 815]
    for idx, (num, verb, q, answer, color) in enumerate(items):
        x, y = xs[idx % 3], ys[idx // 3]
        panel(d, (x, y, x + 710, y + 350), fill=WHITE, outline=INK, width=3, shadow=True)
        d.rectangle((x, y, x + 16, y + 350), fill=color)
        draw_text(d, num, (x + 42, y + 28, x + 180, y + 110), size=55, color=color, bold=True, family="serif")
        draw_text(d, verb, (x + 200, y + 48, x + 650, y + 92), size=22, color=color, bold=True)
        draw_text(d, q, (x + 42, y + 140, x + 650, y + 210), size=34, color=INK, bold=True, family="serif")
        draw_text(d, answer, (x + 42, y + 245, x + 650, y + 320), size=25, color=GRAY)
    return image


def slide_03():
    image, d = base_slide(3, "Classless addressing", "CIDR: a flexible address boundary", "The /n prefix says exactly how many leading bits identify the network.", GREEN)
    # Classful vs CIDR
    panel(d, (110, 350, 825, 670), fill="#FFF8F7", outline=CORAL, width=4)
    tag(d, "Legacy", (140, 380, 285, 425), fill=CORAL, size=19)
    draw_text(d, "Classful", (140, 470, 720, 530), size=42, color=INK, bold=True, family="serif")
    draw_text(d, "Fixed A / B / C boundaries\nAddress space often wasted\nRoute tables could grow quickly",
              (145, 555, 750, 650), size=25, color=GRAY, spacing=9)
    arrow(d, (850, 510), (1030, 510), color=INK, width=5, head=22)
    panel(d, (1060, 350, 1780, 670), fill="#F4FBF8", outline=GREEN, width=4)
    tag(d, "Modern", (1090, 380, 1260, 425), fill=GREEN, size=19)
    draw_text(d, "CIDR", (1090, 470, 1620, 530), size=42, color=INK, bold=True, family="serif")
    draw_text(d, "Any prefix length /0–/32\nRight-sized subnets with VLSM\nRoutes can be aggregated",
              (1095, 555, 1700, 650), size=25, color=GRAY, spacing=9)
    # Formula panel
    panel(d, (1850, 350, 2450, 670), fill=BLUE_PALE, outline=BLUE, width=4)
    draw_text(d, "CIDR FORMULA", (1890, 390, 2390, 430), size=21, color=BLUE, bold=True)
    draw_text(d, "2^(32 − n)", (1890, 485, 2390, 555), size=45, color=INK, bold=True, family="mono", align="center")
    draw_text(d, "total IPv4 addresses", (1890, 585, 2390, 625), size=24, color=GRAY, align="center")

    # Address anatomy
    draw_text(d, "ADDRESS ANATOMY", (110, 760, 620, 800), size=22, color=GREEN, bold=True)
    draw_text(d, "203.0.113.77/27", (110, 825, 960, 900), size=50, color=INK, bold=True, family="mono")
    x0, y0 = 110, 970
    widths = [340, 340, 340, 195, 325]
    labels = ["203", "0", "113", "010", "01101"]
    fills = [GREEN_PALE, GREEN_PALE, GREEN_PALE, GREEN_PALE, YELLOW_PALE]
    x = x0
    for w, label_value, fill in zip(widths, labels, fills):
        d.rectangle((x, y0, x + w, y0 + 100), fill=fill, outline=INK, width=3)
        draw_text(d, label_value, (x, y0, x + w, y0 + 100), size=31, color=INK, bold=True, family="mono",
                  align="center", valign="middle", spacing=0)
        x += w
    draw_text(d, "27 NETWORK BITS", (110, 1100, 1325, 1140), size=23, color=GREEN, bold=True, align="center")
    draw_text(d, "5 HOST BITS", (1325, 1100, 1650, 1140), size=23, color="#A57600", bold=True, align="center")
    rule(d, 110, 1160, 1650, 1160, color=INK, width=3)
    # Prefix quick reference
    panel(d, (1770, 790, 2450, 1245), fill=WHITE, outline=INK, width=3)
    draw_text(d, "QUICK REFERENCE", (1810, 830, 2380, 875), size=22, color=BLUE, bold=True)
    refs = [("/24", "256 addresses", "254 traditional hosts"), ("/26", "64 addresses", "62 traditional hosts"),
            ("/28", "16 addresses", "14 traditional hosts"), ("/30", "4 addresses", "2 traditional hosts")]
    yy = 905
    for prefix, total, usable in refs:
        tag(d, prefix, (1810, yy, 1935, yy + 48), fill=BLUE, size=20)
        draw_text(d, total, (1970, yy + 2, 2200, yy + 48), size=23, color=INK, bold=True)
        draw_text(d, usable, (2200, yy + 5, 2400, yy + 48), size=19, color=GRAY, align="right")
        yy += 78
    return image


def slide_04():
    image, d = base_slide(4, "CIDR worked example", "Decode 192.168.10.70/26", "A repeatable four-step method gives the network, host range, and broadcast.", BLUE)
    steps = [
        ("1", "HOST BITS", "/26 → 32 − 26 = 6", BLUE),
        ("2", "BLOCK SIZE", "2⁶ = 64 addresses", GREEN),
        ("3", "FIND BLOCK", "70 lies in 64–127", CORAL),
        ("4", "WRITE RESULTS", "network .64 • broadcast .127", PURPLE),
    ]
    y = 350
    for n, head, value, color in steps:
        numbered_badge(d, n, (155, y + 42), fill=color)
        draw_text(d, head, (210, y + 3, 560, y + 38), size=20, color=color, bold=True)
        draw_text(d, value, (210, y + 45, 830, y + 92), size=28, color=INK, bold=True)
        if n != "4":
            rule(d, 155, y + 115, 830, y + 115, color=LIGHT, width=2)
        y += 150

    # Binary view
    panel(d, (900, 350, 2450, 735), fill=PALE, outline=INK, width=3)
    draw_text(d, "LAST OCTET IN BINARY", (950, 390, 1750, 430), size=22, color=BLUE, bold=True)
    bits = list("01000110")
    bx, by = 950, 500
    for i, bit in enumerate(bits):
        fill = BLUE_PALE if i < 2 else YELLOW_PALE
        d.rectangle((bx + i * 150, by, bx + i * 150 + 128, by + 105), fill=fill, outline=INK, width=3)
        draw_text(d, bit, (bx + i * 150, by, bx + i * 150 + 128, by + 105), size=42, color=INK,
                  bold=True, family="mono", align="center", valign="middle", spacing=0)
    draw_text(d, "network portion of final octet", (950, 630, 1250, 675), size=21, color=BLUE, bold=True, align="center")
    draw_text(d, "host portion", (1250, 630, 2150, 675), size=21, color="#A57600", bold=True, align="center")

    # Address ruler
    draw_text(d, "WHICH /26 BLOCK CONTAINS 70?", (900, 805, 1800, 850), size=22, color=CORAL, bold=True)
    ruler_x, ruler_y, ruler_w = 900, 900, 1550
    blocks = [("0–63", BLUE_PALE, BLUE), ("64–127", CORAL_PALE, CORAL), ("128–191", GREEN_PALE, GREEN), ("192–255", PURPLE_PALE, PURPLE)]
    bw = ruler_w // 4
    for i, (label_value, fill, color) in enumerate(blocks):
        x = ruler_x + i * bw
        d.rectangle((x, ruler_y, x + bw, ruler_y + 120), fill=fill, outline=INK, width=3)
        draw_text(d, label_value, (x, ruler_y, x + bw, ruler_y + 120), size=27, color=color, bold=True,
                  align="center", valign="middle", spacing=0)
    marker_x = ruler_x + bw + (70 - 64) / 64 * bw
    arrow(d, (marker_x, 1090), (marker_x, 1032), color=CORAL, width=5, head=18)
    draw_text(d, "70", (marker_x - 55, 1100, marker_x + 55, 1140), size=24, color=CORAL, bold=True, align="center")

    panel(d, (900, 1195, 2450, 1315), fill=INK, outline=INK, width=2)
    draw_text(d, "NETWORK  192.168.10.64", (930, 1228, 1450, 1275), size=25, color=WHITE, bold=True, family="mono")
    draw_text(d, "USABLE  .65–.126", (1490, 1228, 1900, 1275), size=25, color=WHITE, bold=True, family="mono")
    draw_text(d, "BROADCAST  .127", (1930, 1228, 2420, 1275), size=25, color=YELLOW, bold=True, family="mono", align="right")
    return image


def slide_05():
    image, d = base_slide(5, "Classless design", "VLSM: give each group the right-sized subnet", "Allocate largest needs first, preserve aligned boundaries, then summarize upstream.", GREEN)
    parent = (110, 360, 2450, 510)
    panel(d, parent, fill=INK, outline=INK, width=3)
    draw_text(d, "PARENT BLOCK", (150, 392, 450, 435), size=21, color=GREEN, bold=True)
    draw_text(d, "10.20.0.0/24", (470, 382, 1100, 450), size=40, color=WHITE, bold=True, family="mono")
    draw_text(d, "256 total addresses • one route can describe the whole site", (1260, 397, 2380, 445), size=26, color="#DCE3EF", align="right")

    # Address ruler scaled by 256 addresses.
    ruler = [(0, 128, "SALES • 100 hosts", "/25", BLUE, BLUE_PALE),
             (128, 192, "ENGINEERING • 50", "/26", GREEN, GREEN_PALE),
             (192, 224, "IoT • 25", "/27", CORAL, CORAL_PALE),
             (224, 256, "SPARE", "/27", PURPLE, PURPLE_PALE)]
    rx, ry, rw, rh = 110, 620, 2340, 220
    for start, end, label_value, prefix, color, fill in ruler:
        x1 = rx + rw * start / 256
        x2 = rx + rw * end / 256
        d.rectangle((x1, ry, x2, ry + rh), fill=fill, outline=INK, width=3)
        draw_text(d, prefix, (x1 + 10, ry + 24, x2 - 10, ry + 75), size=31, color=color, bold=True, family="mono", align="center")
        draw_text(d, label_value, (x1 + 10, ry + 95, x2 - 10, ry + 145), size=22, color=INK, bold=True, align="center")
        draw_text(d, f"{start}–{end - 1}", (x1 + 10, ry + 165, x2 - 10, ry + 205), size=19, color=GRAY, family="mono", align="center")
    draw_text(d, ".0", (100, 850, 200, 890), size=20, color=GRAY, family="mono")
    draw_text(d, ".255", (2310, 850, 2450, 890), size=20, color=GRAY, family="mono", align="right")

    # Allocations and routing ideas.
    panels = [
        (110, 950, 830, 1285, BLUE, "01 • ALLOCATE", "10.20.0.0/25", "usable .1–.126"),
        (890, 950, 1610, 1285, GREEN, "02 • CONTINUE", "10.20.0.128/26", "usable .129–.190"),
        (1670, 950, 2450, 1285, CORAL, "03 • MOST SPECIFIC WINS", "10.20.0.192/27", "a /27 route beats a matching /24"),
    ]
    for x1, y1, x2, y2, color, head, value, note in panels:
        panel(d, (x1, y1, x2, y2), fill=WHITE, outline=color, width=4)
        draw_text(d, head, (x1 + 30, y1 + 30, x2 - 30, y1 + 70), size=20, color=color, bold=True)
        draw_text(d, value, (x1 + 30, y1 + 115, x2 - 30, y1 + 175), size=31, color=INK, bold=True, family="mono")
        draw_text(d, note, (x1 + 30, y1 + 220, x2 - 30, y2 - 30), size=24, color=GRAY)
    return image


def slide_06():
    image, d = base_slide(6, "Network Address Translation", "NAT/PAT: rewrite at the network edge", "Many private hosts can share public IPv4 reachability through address-and-port mappings.", CORAL)
    # Zones
    panel(d, (110, 355, 730, 900), fill=GREEN_PALE, outline=GREEN, width=4)
    tag(d, "Private LAN", (145, 390, 370, 438), fill=GREEN, size=20)
    draw_text(d, "Client", (155, 520, 420, 565), size=30, color=INK, bold=True, family="serif")
    draw_text(d, "10.0.0.15", (155, 600, 600, 650), size=31, color=INK, bold=True, family="mono")
    draw_text(d, "source port 51514", (155, 670, 600, 715), size=24, color=GRAY, family="mono")
    switch(d, (470, 750, 650, 835), "LAN", GREEN)

    panel(d, (940, 355, 1590, 900), fill=CORAL_PALE, outline=CORAL, width=5)
    tag(d, "NAT Gateway", (990, 390, 1270, 438), fill=CORAL, size=20)
    router(d, (1265, 585), "NAT", CORAL, radius=90)
    draw_text(d, "inside: 10.0.0.1", (1010, 735, 1520, 775), size=24, color=INK, family="mono", align="center")
    draw_text(d, "outside: 203.0.113.8", (1010, 795, 1520, 840), size=24, color=INK, family="mono", align="center")

    panel(d, (1800, 355, 2450, 900), fill=BLUE_PALE, outline=BLUE, width=4)
    tag(d, "Public Internet", (1840, 390, 2150, 438), fill=BLUE, size=20)
    draw_text(d, "Web server", (1850, 520, 2280, 565), size=30, color=INK, bold=True, family="serif")
    draw_text(d, "198.51.100.20", (1850, 600, 2380, 650), size=31, color=INK, bold=True, family="mono")
    draw_text(d, "HTTPS port 443", (1850, 670, 2340, 715), size=24, color=GRAY, family="mono")

    arrow(d, (730, 550), (925, 550), color=CORAL, width=8, head=24)
    packet(d, (755, 490, 895, 535), "1 • OUT", CORAL)
    arrow(d, (1590, 550), (1785, 550), color=CORAL, width=8, head=24)
    arrow(d, (1785, 720), (1590, 720), color=BLUE, width=7, head=24)
    arrow(d, (925, 720), (730, 720), color=BLUE, width=7, head=24)
    packet(d, (1630, 655, 1765, 700), "2 • IN", BLUE)

    # Mapping table
    panel(d, (110, 1000, 2450, 1285), fill=WHITE, outline=INK, width=3)
    draw_text(d, "PAT TRANSLATION TABLE", (150, 1035, 620, 1075), size=21, color=CORAL, bold=True)
    headers = ["PROTO", "INSIDE LOCAL", "INSIDE GLOBAL", "DESTINATION"]
    values = ["TCP", "10.0.0.15:51514", "203.0.113.8:40001", "198.51.100.20:443"]
    xs = [150, 480, 1050, 1690, 2400]
    for i, head in enumerate(headers):
        draw_text(d, head, (xs[i], 1115, xs[i + 1] - 20, 1155), size=19, color=GRAY, bold=True)
        draw_text(d, values[i], (xs[i], 1190, xs[i + 1] - 20, 1245), size=25, color=INK, bold=True, family="mono")
        if i:
            rule(d, xs[i] - 30, 1100, xs[i] - 30, 1260, color=LIGHT, width=2)
    return image


def slide_07():
    image, d = base_slide(7, "NAT patterns", "Three NAT types—and what NAT does not do", "Choose the mapping model deliberately; address translation changes end-to-end behavior.", CORAL)
    types = [
        (110, BLUE, "STATIC NAT", "1 private to 1 public", "Permanent mapping", "Useful for a reachable internal service"),
        (910, GREEN, "DYNAMIC NAT", "many to public pool", "Temporary mapping", "A public address is borrowed while active"),
        (1710, CORAL, "PAT / NAPT", "many sessions to 1 IP", "Ports identify flows", "The most common home/enterprise form"),
    ]
    for x, color, head, mapping, mode, use in types:
        panel(d, (x, 360, x + 720, 850), fill=WHITE, outline=color, width=4, shadow=True)
        tag(d, head, (x + 30, 390, x + 330, 442), fill=color, size=20)
        draw_text(d, mapping, (x + 40, 510, x + 670, 575), size=35, color=INK, bold=True, family="serif", align="center")
        # little mapping illustration
        for j in range(3):
            d.ellipse((x + 70, 640 + j * 42, x + 92, 662 + j * 42), fill=color)
        arrow(d, (x + 130, 690), (x + 515, 690), color=color, width=5, head=18)
        d.rectangle((x + 540, 640, x + 650, 742), fill=WHITE, outline=color, width=4)
        draw_text(d, "IP", (x + 540, 640, x + 650, 742), size=24, color=INK, bold=True, align="center", valign="middle")
        draw_text(d, mode, (x + 40, 770, x + 680, 810), size=23, color=color, bold=True, align="center")
        draw_text(d, use, (x + 60, 815, x + 660, 840), size=19, color=GRAY, align="center")

    # Benefits and limits
    panel(d, (110, 955, 1240, 1285), fill=GREEN_PALE, outline=GREEN, width=4)
    draw_text(d, "WHAT NAT HELPS", (150, 995, 600, 1035), size=21, color=GREEN, bold=True)
    helps = ["Conserves scarce public IPv4 addresses", "Hides internal addressing from external routes", "Lets an organization renumber internally"]
    yy = 1080
    for item in helps:
        d.rectangle((155, yy + 8, 171, yy + 24), fill=GREEN)
        draw_text(d, item, (195, yy, 1160, yy + 50), size=24, color=INK)
        yy += 67
    panel(d, (1320, 955, 2450, 1285), fill=CORAL_PALE, outline=CORAL, width=4)
    draw_text(d, "WHAT NAT DOES NOT GUARANTEE", (1360, 995, 2350, 1035), size=21, color=CORAL, bold=True)
    limits = ["It is not a firewall or encryption", "Unsolicited inbound traffic needs a mapping", "Some apps dislike embedded IP addresses"]
    yy = 1080
    for item in limits:
        d.rectangle((1365, yy + 8, 1381, yy + 24), fill=CORAL)
        draw_text(d, item, (1405, yy, 2370, yy + 50), size=24, color=INK)
        yy += 67
    return image


def slide_08():
    image, d = base_slide(8, "IPv4 to IPv6", "Transition toolbox: match the constraint", "Dual stack and tunnels aid transition; SIIT, NAT64, and 464XLAT perform translation.", PURPLE)
    rows = [
        ("NATIVE / DUAL STACK", "speak both", "No translation", "Both endpoints and networks support IPv6", BLUE, BLUE_PALE),
        ("TUNNEL", "wrap", "Encapsulation", "Carry IPv6 islands across an IPv4 core", YELLOW, YELLOW_PALE),
        ("SIIT", "map", "Stateless translation", "Deterministic IPv4↔IPv6 header/address mapping", GREEN, GREEN_PALE),
        ("NAT64 + DNS64", "state", "Stateful translation", "IPv6-only clients reach IPv4-only servers", PURPLE, PURPLE_PALE),
        ("464XLAT", "twice", "CLAT + PLAT", "IPv4-only apps work over IPv6-only access", CORAL, CORAL_PALE),
    ]
    col_x = [110, 660, 1120, 1720, 2450]
    headers = ["STRATEGY", "PACKET ACTION", "STATE MODEL", "BEST WHEN"]
    for i, head in enumerate(headers):
        draw_text(d, head, (col_x[i] + 10, 355, col_x[i + 1] - 20, 400), size=19, color=GRAY, bold=True)
    y = 430
    for strategy, action, state, best, color, fill in rows:
        d.rectangle((110, y, 2450, y + 155), fill=fill)
        d.rectangle((110, y, 126, y + 155), fill=color)
        draw_text(d, strategy, (150, y + 38, 620, y + 105), size=27, color=INK, bold=True, family="serif")
        tag(d, action, (700, y + 49, 1000, y + 103), fill=color, size=19)
        draw_text(d, state, (1150, y + 45, 1660, y + 105), size=25, color=INK, bold=True)
        draw_text(d, best, (1750, y + 34, 2400, y + 118), size=24, color=INK)
        y += 174
    panel(d, (110, 1320 - 92, 2450, 1320), fill=INK, outline=INK, width=2)
    draw_text(d, "MEMORY HOOK", (145, 1249, 430, 1295), size=20, color=YELLOW, bold=True)
    draw_text(d, "Dual stack speaks • Tunnel wraps • Translation rewrites", (480, 1244, 2350, 1300), size=28, color=WHITE, bold=True, align="center")
    return image


def slide_09():
    image, d = base_slide(9, "Worked translation flow", "NAT64 + DNS64 in five moves", "An IPv6-only client reaches an IPv4-only service without the server changing.", PURPLE)
    nodes = [
        (150, BLUE, "1", "IPv6 client", "2001:db8::10"),
        (610, PURPLE, "2", "DNS64", "query AAAA"),
        (1070, YELLOW, "3", "Synthesized target", "64:ff9b::c000:222"),
        (1590, GREEN, "4", "NAT64", "session + rewrite"),
        (2070, CORAL, "5", "IPv4 server", "192.0.2.34"),
    ]
    for i in range(len(nodes) - 1):
        arrow(d, (nodes[i][0] + 300, 590), (nodes[i + 1][0] - 40, 590), color=INK, width=5, head=20)
    for x, color, n, head, sub in nodes:
        panel(d, (x, 420, x + 340, 760), fill=WHITE, outline=color, width=5)
        numbered_badge(d, n, (x + 55, 475), fill=color, diameter=58, color=INK if color == YELLOW else WHITE)
        draw_text(d, head, (x + 28, 555, x + 312, 620), size=29, color=INK, bold=True, family="serif", align="center")
        draw_text(d, sub, (x + 20, 670, x + 320, 730), size=20, color=GRAY, family="mono", align="center")

    explanations = [
        (150, "Ask DNS for an AAAA record"),
        (610, "If no AAAA exists, use the A record"),
        (1070, "Embed 192.0.2.34 in the NAT64 prefix"),
        (1590, "Convert IPv6 ↔ IPv4 and track the flow"),
        (2070, "Reply follows the mapping in reverse"),
    ]
    for x, value in explanations:
        draw_text(d, value, (x, 800, x + 340, 900), size=22, color=INK, bold=True, align="center")

    panel(d, (110, 1020, 2450, 1275), fill=PALE, outline=INK, width=3)
    draw_text(d, "ADDRESS SYNTHESIS", (150, 1055, 540, 1095), size=20, color=PURPLE, bold=True)
    draw_text(d, "64:ff9b::/96", (150, 1140, 630, 1195), size=30, color=PURPLE, bold=True, family="mono")
    draw_text(d, "+", (665, 1140, 730, 1195), size=32, color=INK, bold=True, align="center")
    draw_text(d, "192.0.2.34", (760, 1140, 1170, 1195), size=30, color=CORAL, bold=True, family="mono")
    arrow(d, (1200, 1164), (1430, 1164), color=INK, width=5, head=20)
    draw_text(d, "64:ff9b::c000:222", (1480, 1130, 2360, 1205), size=34, color=INK, bold=True, family="mono", align="center")
    draw_text(d, "Watch for IPv4 literals and protocols that embed addresses in payloads.", (150, 1230, 2380, 1260), size=21, color=GRAY, align="center")
    return image


def slide_10():
    image, d = base_slide(10, "Packet network topology", "The map offers paths; routing chooses one", "Edge, access, and core layers connect end systems through shared packet links.", BLUE)
    # Layer bands
    bands = [(110, 360, 620, 1030, "EDGE", BLUE, BLUE_PALE), (650, 360, 1080, 1030, "ACCESS", GREEN, GREEN_PALE),
             (1110, 360, 1980, 1030, "CORE MESH", PURPLE, PURPLE_PALE), (2010, 360, 2450, 1030, "EDGE", CORAL, CORAL_PALE)]
    for x1, y1, x2, y2, name, color, fill in bands:
        d.rectangle((x1, y1, x2, y2), fill=fill)
        tag(d, name, (x1 + 25, y1 + 25, min(x2 - 25, x1 + 230), y1 + 72), fill=color, size=18)
    # Links behind nodes
    links = [((470, 675), (800, 675)), ((900, 675), (1250, 510)), ((900, 675), (1250, 800)),
             ((1340, 510), (1550, 650)), ((1340, 800), (1550, 650)), ((1640, 650), (1830, 490)),
             ((1640, 650), (1830, 825)), ((1920, 490), (2180, 675)), ((1920, 825), (2180, 675))]
    for a, b in links:
        rule(d, *a, *b, color="#A8B2C3", width=5)
    selected = [((470, 675), (800, 675)), ((900, 675), (1250, 510)), ((1340, 510), (1550, 650)),
                ((1640, 650), (1830, 490)), ((1920, 490), (2180, 675))]
    for a, b in selected:
        arrow(d, a, b, color=CORAL, width=9, head=22)
    # nodes
    panel(d, (190, 585, 470, 765), fill=WHITE, outline=BLUE, width=5)
    draw_text(d, "HOST A", (190, 610, 470, 660), size=25, color=INK, bold=True, align="center")
    draw_text(d, "10.20.0.34", (190, 690, 470, 735), size=20, color=GRAY, family="mono", align="center")
    switch(d, (800, 620, 900, 730), "SW", GREEN)
    router(d, (1295, 510), "R1", PURPLE)
    router(d, (1295, 800), "R2", PURPLE)
    router(d, (1595, 650), "R3", PURPLE)
    router(d, (1875, 490), "R4", PURPLE)
    router(d, (1875, 825), "R5", PURPLE)
    panel(d, (2180, 585, 2405, 765), fill=WHITE, outline=CORAL, width=5)
    draw_text(d, "SERVER", (2180, 610, 2405, 660), size=25, color=INK, bold=True, align="center")
    draw_text(d, "destination", (2180, 690, 2405, 735), size=20, color=GRAY, align="center")
    packet(d, (1450, 480, 1545, 520), "P", CORAL)

    # Definitions
    terms = [(110, "NODE", "host, switch, router", BLUE), (690, "LINK", "wired or wireless channel", GREEN),
             (1280, "HOP", "one forwarding step", PURPLE), (1830, "PATH", "sequence of links", CORAL)]
    for x, head, desc, color in terms:
        draw_text(d, head, (x, 1110, x + 150, 1150), size=20, color=color, bold=True)
        draw_text(d, desc, (x, 1160, x + 500, 1210), size=23, color=INK)
    panel(d, (110, 1265, 2450, 1335), fill=INK, outline=INK, width=2)
    draw_text(d, "PHYSICAL topology = cables/radio  •  LOGICAL topology = the forwarding path traffic follows",
              (150, 1282, 2410, 1325), size=24, color=WHITE, bold=True, align="center")
    return image


def slide_11():
    image, d = base_slide(11, "Network service models", "Datagrams vs. virtual circuits", "The key difference: decide independently per packet—or establish a logical path first.", PURPLE)
    # Datagram side
    panel(d, (110, 360, 1220, 885), fill=BLUE_PALE, outline=BLUE, width=4)
    tag(d, "Datagram", (145, 390, 380, 442), fill=BLUE, size=21)
    draw_text(d, "NO SETUP", (430, 400, 650, 435), size=19, color=BLUE, bold=True)
    pts = {"S": (230, 650), "A": (520, 500), "B": (520, 790), "C": (840, 645), "D": (1090, 645)}
    dgram_links = [("S", "A"), ("S", "B"), ("A", "C"), ("B", "C"), ("C", "D")]
    for a, b in dgram_links:
        rule(d, *pts[a], *pts[b], color="#A8B2C3", width=4)
    for a, b in [("S", "A"), ("A", "C"), ("C", "D")]:
        arrow(d, pts[a], pts[b], color=BLUE, width=7, head=20)
    for a, b in [("S", "B"), ("B", "C"), ("C", "D")]:
        arrow(d, (pts[a][0], pts[a][1] + 10), (pts[b][0], pts[b][1] + 10), color=CORAL, width=6, head=18)
    for name, p in pts.items():
        router(d, p, name, BLUE if name in ("S", "D") else INK, radius=38)
    packet(d, (640, 505, 735, 545), "P1", BLUE)
    packet(d, (640, 755, 735, 795), "P2", CORAL)
    draw_text(d, "Each packet carries the destination; routes may differ and arrival can reorder.",
              (160, 820, 1160, 865), size=22, color=INK, bold=True, align="center")

    # VC side
    panel(d, (1340, 360, 2450, 885), fill=PURPLE_PALE, outline=PURPLE, width=4)
    tag(d, "Virtual circuit", (1375, 390, 1695, 442), fill=PURPLE, size=21)
    draw_text(d, "SETUP → TRANSFER → TEARDOWN", (1750, 400, 2390, 435), size=19, color=PURPLE, bold=True, align="right")
    route_pts = [(1450, 650, "S"), (1710, 650, "R1"), (1980, 650, "R2"), (2320, 650, "D")]
    for i in range(len(route_pts) - 1):
        arrow(d, (route_pts[i][0] + 42, 650), (route_pts[i + 1][0] - 42, 650), color=PURPLE, width=9, head=22)
    labels = ["VC 42", "VC 17", "VC 8"]
    for i, (x, y, name) in enumerate(route_pts):
        router(d, (x, y), name, PURPLE, radius=42)
        if i < len(labels):
            tag(d, labels[i], (x + 65, 570, x + 180, 610), fill=WHITE, color=PURPLE, size=16)
    draw_text(d, "Switches keep per-circuit state; packets usually follow one path and remain ordered.",
              (1390, 820, 2400, 865), size=22, color=INK, bold=True, align="center")

    # Comparison table
    table_y = 965
    headers = ["QUESTION", "DATAGRAM", "VIRTUAL CIRCUIT"]
    col = [110, 720, 1540, 2450]
    for i, h in enumerate(headers):
        d.rectangle((col[i], table_y, col[i + 1], table_y + 65), fill=INK if i == 0 else (BLUE if i == 1 else PURPLE))
        draw_text(d, h, (col[i] + 20, table_y, col[i + 1] - 20, table_y + 65), size=20, color=WHITE, bold=True,
                  align="center", valign="middle", spacing=0)
    rows = [("Setup", "none", "required"), ("Packet identity", "destination address", "short label / VC ID"),
            ("Failure", "later packets can reroute", "circuit may be rebuilt"), ("Examples", "IP", "ATM / Frame Relay / MPLS-like")]
    y = table_y + 65
    for r, row in enumerate(rows):
        fill = WHITE if r % 2 == 0 else PALE
        for i, value in enumerate(row):
            d.rectangle((col[i], y, col[i + 1], y + 67), fill=fill, outline=LIGHT, width=2)
            draw_text(d, value, (col[i] + 20, y, col[i + 1] - 20, y + 67), size=21, color=INK, bold=i == 0,
                      align="center", valign="middle", spacing=0)
        y += 67
    return image


def slide_12():
    image, d = base_slide(12, "Forwarding hardware", "Inside a switch/router", "The data plane moves packets quickly; the control plane decides what the tables should contain.", CORAL)
    # Control plane
    panel(d, (760, 350, 1800, 540), fill=YELLOW_PALE, outline=YELLOW, width=4)
    tag(d, "Control plane", (800, 385, 1055, 433), fill=YELLOW, color=INK, size=20)
    draw_text(d, "CPU / routing protocols / management", (1100, 392, 1740, 438), size=27, color=INK, bold=True, align="center")
    draw_text(d, "builds MAC, ARP/ND, and forwarding tables", (820, 470, 1740, 510), size=23, color=GRAY, align="center")
    arrow(d, (1280, 545), (1280, 650), color=YELLOW, width=6, head=22)

    # Data plane pipeline
    blocks = [
        (110, 660, 590, 1050, BLUE, "INPUT PORT", ["PHY + MAC", "frame check", "header parse"]),
        (650, 660, 1130, 1050, GREEN, "LOOKUP", ["MAC table", "longest prefix", "ACL / policy"]),
        (1190, 660, 1670, 1050, PURPLE, "SWITCH FABRIC", ["memory", "shared bus", "crossbar"]),
        (1730, 660, 2450, 1050, CORAL, "OUTPUT PORT", ["buffer / queue", "scheduler", "PHY transmit"]),
    ]
    for i, (x1, y1, x2, y2, color, head, items) in enumerate(blocks):
        panel(d, (x1, y1, x2, y2), fill=WHITE, outline=color, width=5)
        tag(d, head, (x1 + 25, y1 + 25, x2 - 25, y1 + 78), fill=color, size=19)
        yy = y1 + 145
        for item in items:
            d.rectangle((x1 + 55, yy + 8, x1 + 71, yy + 24), fill=color)
            draw_text(d, item, (x1 + 95, yy, x2 - 35, yy + 48), size=25, color=INK, bold=True)
            yy += 68
        if i < len(blocks) - 1:
            arrow(d, (x2 + 10, (y1 + y2) // 2), (blocks[i + 1][0] - 15, (y1 + y2) // 2), color=INK, width=5, head=20)

    # Compare L2/L3 and final recap
    panel(d, (110, 1130, 1190, 1315), fill=BLUE_PALE, outline=BLUE, width=3)
    draw_text(d, "L2 SWITCH", (150, 1160, 390, 1200), size=21, color=BLUE, bold=True)
    draw_text(d, "destination MAC → MAC table → output port", (420, 1150, 1135, 1210), size=25, color=INK, bold=True)
    draw_text(d, "forwards frames inside a LAN", (420, 1230, 1135, 1275), size=22, color=GRAY)
    panel(d, (1260, 1130, 2450, 1315), fill=CORAL_PALE, outline=CORAL, width=3)
    draw_text(d, "L3 ROUTER", (1300, 1160, 1560, 1200), size=21, color=CORAL, bold=True)
    draw_text(d, "destination IP → prefix table → next hop", (1600, 1150, 2395, 1210), size=25, color=INK, bold=True)
    draw_text(d, "decrements TTL and builds a new link frame", (1600, 1230, 2395, 1275), size=22, color=GRAY)
    draw_text(d, "PLAN ADDRESS → SWITCH FRAME → ROUTE PACKET → TRANSLATE IF NEEDED → DELIVER",
              (170, 1333, 2390, 1365), size=20, color=CORAL, bold=True, align="center")
    return image


def build():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    slides = [slide_01(), slide_02(), slide_03(), slide_04(), slide_05(), slide_06(),
              slide_07(), slide_08(), slide_09(), slide_10(), slide_11(), slide_12()]
    pngs = []
    for idx, image in enumerate(slides, 1):
        path = ASSET_DIR / f"slide_{idx:02d}.png"
        image.save(path, format="PNG", optimize=True, compress_level=7, dpi=(192, 192))
        pngs.append(path)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for idx, path in enumerate(pngs, 1):
        slide = prs.slides.add_slide(blank)
        picture = slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
        picture.name = f"Phone-safe slide {idx:02d}"
    prs.core_properties.title = "Computer Networks: From Prefixes to Packets"
    prs.core_properties.subject = "CIDR, VLSM, NAT, IPv4/IPv6 transition, topology, datagrams, virtual circuits, switches, and routers"
    prs.core_properties.author = "OpenAI"
    prs.core_properties.comments = "12-slide white-background mobile-safe visual guide"
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
