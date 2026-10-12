"""Flatten a shape-heavy PowerPoint into high-resolution slide images.

Some mobile PowerPoint viewers render connectors, custom shapes, and substituted
fonts inconsistently. This utility redraws each slide at 2560x1440 and places one
PNG on each slide, making the result visually stable across phone viewers.
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN


WIDTH = 2560
HEIGHT = 1440
EMU_PER_INCH = 914400
PT_PER_INCH = 72

FONT_DIR = Path("/mnt/c/Windows/Fonts")
REGULAR_FONT = FONT_DIR / "calibri.ttf"
BOLD_FONT = FONT_DIR / "calibrib.ttf"
MONO_FONT = FONT_DIR / "consola.ttf"
MONO_BOLD_FONT = FONT_DIR / "consolab.ttf"

if not REGULAR_FONT.exists():
    FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")
    REGULAR_FONT = FONT_DIR / "DejaVuSans.ttf"
    BOLD_FONT = FONT_DIR / "DejaVuSans-Bold.ttf"
    MONO_FONT = FONT_DIR / "DejaVuSansMono.ttf"
    MONO_BOLD_FONT = FONT_DIR / "DejaVuSansMono-Bold.ttf"


def color(value, fallback=(0, 0, 0)):
    try:
        raw = str(value.rgb)
        return tuple(int(raw[i : i + 2], 16) for i in (0, 2, 4))
    except Exception:
        return fallback


def attr_is_true(element, attribute):
    return element.get(attribute) in {"1", "true", "True"}


class SlideRenderer:
    def __init__(self, presentation: Presentation):
        self.prs = presentation
        self.sx = WIDTH / presentation.slide_width
        self.sy = HEIGHT / presentation.slide_height
        self.pixels_per_inch = WIDTH / (presentation.slide_width / EMU_PER_INCH)

    def box(self, shape):
        return (
            round(shape.left * self.sx),
            round(shape.top * self.sy),
            round((shape.left + shape.width) * self.sx),
            round((shape.top + shape.height) * self.sy),
        )

    def font(self, run, paragraph, point_size):
        name = (run.font.name or "").lower()
        bold = bool(run.font.bold)
        mono = "consol" in name or "mono" in name
        path = MONO_BOLD_FONT if mono and bold else MONO_FONT if mono else BOLD_FONT if bold else REGULAR_FONT
        pixel_size = max(12, round(point_size * self.pixels_per_inch / PT_PER_INCH))
        return ImageFont.truetype(str(path), pixel_size)

    @staticmethod
    def wrap(draw, text, font, width):
        output = []
        for hard_line in text.split("\n"):
            if not hard_line:
                output.append("")
                continue
            words = hard_line.split(" ")
            current = ""
            for word in words:
                candidate = word if not current else f"{current} {word}"
                if not current or draw.textlength(candidate, font=font) <= width:
                    current = candidate
                else:
                    output.append(current)
                    current = word
            output.append(current)
        return output

    @staticmethod
    def dashed_line(draw, start, end, fill, width, dash=15, gap=10):
        x1, y1 = start
        x2, y2 = end
        length = math.hypot(x2 - x1, y2 - y1)
        if length == 0:
            return
        ux = (x2 - x1) / length
        uy = (y2 - y1) / length
        pos = 0.0
        while pos < length:
            stop = min(pos + dash, length)
            draw.line(
                (x1 + ux * pos, y1 + uy * pos, x1 + ux * stop, y1 + uy * stop),
                fill=fill,
                width=width,
            )
            pos += dash + gap

    def render_line(self, draw, shape):
        x1, y1, x2, y2 = self.box(shape)
        try:
            xfrm = shape._element.spPr.xfrm
            reverse_slope = attr_is_true(xfrm, "flipH") ^ attr_is_true(xfrm, "flipV")
        except Exception:
            reverse_slope = False
        if reverse_slope:
            start, end = (x1, y2), (x2, y1)
        else:
            start, end = (x1, y1), (x2, y2)
        stroke = color(shape.line.color, (120, 130, 150))
        line_width = max(1, round((shape.line.width or 12700) / EMU_PER_INCH * self.pixels_per_inch))
        try:
            dashed = "dash" in str(shape.line.dash_style).lower()
        except Exception:
            dashed = False
        if dashed:
            self.dashed_line(draw, start, end, stroke, line_width)
        else:
            draw.line((*start, *end), fill=stroke, width=line_width)

    def render_auto_shape(self, draw, shape):
        x1, y1, x2, y2 = self.box(shape)
        fill = color(shape.fill.fore_color, (255, 255, 255))
        stroke = color(shape.line.color, fill)
        line_width = max(1, round((shape.line.width or 12700) / EMU_PER_INCH * self.pixels_per_inch))
        kind = shape.auto_shape_type
        if kind == MSO_AUTO_SHAPE_TYPE.OVAL:
            draw.ellipse((x1, y1, x2, y2), fill=fill, outline=stroke, width=line_width)
        elif kind == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
            radius = max(6, min(x2 - x1, y2 - y1) // 6)
            draw.rounded_rectangle((x1, y1, x2, y2), radius=radius, fill=fill, outline=stroke, width=line_width)
        elif kind == MSO_AUTO_SHAPE_TYPE.CHEVRON:
            notch = (x2 - x1) * 0.25
            points = [
                (x1, y1),
                (x2 - notch, y1),
                (x2, (y1 + y2) / 2),
                (x2 - notch, y2),
                (x1, y2),
                (x1 + notch, (y1 + y2) / 2),
            ]
            draw.polygon(points, fill=fill)
        else:
            draw.rectangle((x1, y1, x2, y2), fill=fill, outline=stroke, width=line_width)

    def render_text(self, draw, shape):
        if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
            return
        x1, y1, x2, y2 = self.box(shape)
        tf = shape.text_frame
        left = x1 + round(tf.margin_left * self.sx)
        right = x2 - round(tf.margin_right * self.sx)
        top = y1 + round(tf.margin_top * self.sy)
        bottom = y2 - round(tf.margin_bottom * self.sy)
        width = max(2, right - left)
        height = max(2, bottom - top)

        rendered_lines = []
        for paragraph in tf.paragraphs:
            if not paragraph.runs:
                continue
            run = paragraph.runs[0]
            point_size = float(run.font.size.pt if run.font.size else 12)
            font = self.font(run, paragraph, point_size)
            text_color = color(run.font.color, (23, 32, 51))
            wrapped = self.wrap(draw, paragraph.text, font, width)
            ascent, descent = font.getmetrics()
            line_height = max(1, round((ascent + descent) * 0.92))
            spacing = max(1, round(line_height * 0.10))
            for value in wrapped:
                rendered_lines.append((value, font, text_color, paragraph.alignment, line_height, spacing))

        total_height = sum(item[4] + item[5] for item in rendered_lines)
        if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
            cursor_y = top + max(0, (height - total_height) // 2)
        elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
            cursor_y = top + max(0, height - total_height)
        else:
            cursor_y = top

        for value, font, text_color, alignment, line_height, spacing in rendered_lines:
            text_width = draw.textlength(value, font=font)
            if alignment == PP_ALIGN.CENTER:
                cursor_x = left + (width - text_width) / 2
            elif alignment == PP_ALIGN.RIGHT:
                cursor_x = right - text_width
            else:
                cursor_x = left
            draw.text((round(cursor_x), cursor_y), value, font=font, fill=text_color)
            cursor_y += line_height + spacing

    def render(self, slide):
        background = (247, 248, 252)
        try:
            background = color(slide.background.fill.fore_color, background)
        except Exception:
            pass
        image = Image.new("RGB", (WIDTH, HEIGHT), background)
        draw = ImageDraw.Draw(image)
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                self.render_line(draw, shape)
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                self.render_auto_shape(draw, shape)
            self.render_text(draw, shape)
        return image


def flatten(source: Path, output: Path, image_dir: Path):
    original = Presentation(source)
    renderer = SlideRenderer(original)
    image_dir.mkdir(parents=True, exist_ok=True)

    images = []
    for index, slide in enumerate(original.slides, start=1):
        image = renderer.render(slide)
        image_path = image_dir / f"slide_{index:02d}.png"
        image.save(image_path, "PNG", optimize=True, dpi=(192, 192))
        images.append(image_path)

    flattened = Presentation()
    flattened.slide_width = original.slide_width
    flattened.slide_height = original.slide_height
    blank = flattened.slide_layouts[6]
    for image_path in images:
        slide = flattened.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image_path),
            0,
            0,
            width=flattened.slide_width,
            height=flattened.slide_height,
        )

    flattened.core_properties.title = f"{original.core_properties.title} — Mobile-safe"
    flattened.core_properties.subject = "Phone-compatible flattened presentation"
    flattened.core_properties.author = original.core_properties.author
    flattened.core_properties.comments = "Each slide is a single 2560×1440 PNG for consistent mobile rendering."
    flattened.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--image-dir", type=Path, default=Path("/tmp/network_mobile_slides"))
    args = parser.parse_args()
    flatten(args.source, args.output, args.image_dir)
    print(args.output)


if __name__ == "__main__":
    main()
